"""Import from Excel (app/sheet_import.py): reading sheets, matching their
columns to a database's, tidying values, and importing through each
database's own save code."""
from __future__ import annotations

import io
import json
import re
from datetime import datetime

# tests.base first: it points the app at a throwaway database before app is imported.
from tests.base import AppTestCase, count, flash_text, last_batch, one, row, rows, uniq
from app import sheet_import as si  # noqa: E402


def xlsx(rows: list[list], sheet: str = "Sheet1", extra: dict | None = None, merge: tuple[str, ...] = ()) -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    for cells in merge:
        ws.merge_cells(cells)
    for name, more in (extra or {}).items():
        other = book.create_sheet(name)
        for r in more:
            other.append(r)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class Matching(AppTestCase):
    def fields(self, target_key="plasmids"):
        from app.app import app
        from app.db import SessionLocal
        with app.test_request_context(), SessionLocal() as s:
            return si.target_for(s, target_key).fields

    def test_names_are_compared_loosely(self):
        self.assertEqual(si.norm("Cat. No."), "cat number")
        self.assertEqual(si.norm("Positions"), "position")
        self.assertEqual(si.norm("Plasmid #"), "plasmid number")

    def test_synonyms_and_values_decide(self):
        fields = self.fields()
        headers = ["Plasmid Name", "Vector", "Antibiotic", "Location", "Freezer", "Who made it", "Box"]
        columns = [["pA"], ["pUC19"], ["Amp"], ["A1", "B2", "C10"], ["-80 top"], ["me"], ["Box 1"]]
        got = {headers[i]: key for i, (key, _s, _w) in si.auto_match(headers, columns, fields).items()}
        self.assertEqual(got["Plasmid Name"], "name")
        self.assertEqual(got["Vector"], "backbone")
        self.assertEqual(got["Antibiotic"], "resistance")
        self.assertEqual(got["Location"], "position")          # it holds A1, B2…
        self.assertEqual(got["Freezer"], "location")
        self.assertEqual(got["Box"], "box")

    def test_a_location_of_words_is_a_location(self):
        fields = self.fields()
        got = si.auto_match(["Location"], [["Freezer 2, shelf 3", "Cold room"]], fields)
        self.assertEqual(got[0][0], "location")

    def test_a_position_column_may_mix_a1_and_plain_numbers(self):
        for values in (["A1", "3", "B2", "12", "C4"], ["1", "2", "3", "81"]):
            got = si.auto_match(["Position"], [values], self.fields())
            self.assertEqual(got[0][0], "position", values)

    def test_merged_cells_cost_the_sheet_s_size_not_the_number_of_merges(self):
        import time
        rows = [[str(i)] for i in range(5000)]
        start = time.monotonic()
        si._fill_merged(rows, [(1, 1 + i, 1, 5001 + i) for i in range(20000)])
        self.assertLess(time.monotonic() - start, 1.0)
        rows = [["C1", "a"], ["", "b"], ["", "c"], ["C2", "d"], ["", "e"]]
        si._fill_merged(rows, [(1, 4, 1, 5), (1, 1, 1, 3)])
        self.assertEqual([r[0] for r in rows], ["C1", "C1", "C1", "C2", "C2"])

    def test_near_spellings_match(self):
        got = si.auto_match(["Resistence"], [["Kan"]], self.fields())
        self.assertEqual(got[0][0], "resistance")

    def test_a_name_with_a_unit_or_a_generic_word_is_that_column(self):
        fields = [si.Field("attr_hazard", "Hazard", custom=True), si.Field("attr_weight", "Weight", kind="number"),
                  si.Field("cage_id", "Cage", ("cage",))]
        fields.append(si.Field("mouse_id", "Mouse ID", ("mouse",)))
        got = si.auto_match(["Hazard class", "Weight (g)", "Cage colour", "Mouse age"],
                            [["toxic"], ["4.5"], ["blue"], ["8"]], fields)
        self.assertEqual({i: k for i, (k, _s, _w) in got.items()}, {0: "attr_hazard", 1: "attr_weight"})

    def test_each_database_column_takes_one_sheet_column(self):
        got = si.auto_match(["Name", "Plasmid name"], [["a"], ["b"]], self.fields())
        self.assertEqual([k for k, _s, _w in got.values()].count("name"), 1)


class Tidying(AppTestCase):
    def test_dates_day_or_month_first_per_column(self):
        out, notes = si.tidy_dates(["14/03/2026", "03/04/2026", ""])
        self.assertEqual(out[:2], ["2026-03-14", "2026-04-03"])
        out, notes = si.tidy_dates(["03/14/2026", "04/03/2026"])
        self.assertEqual(out, ["2026-03-14", "2026-04-03"])
        out, notes = si.tidy_dates(["03/04/2026"])
        self.assertEqual(out, ["2026-03-04"])
        self.assertIn("month first", " ".join(notes))

    def test_dashed_two_digit_years_are_day_or_month_first_not_year_first(self):
        out, _notes = si.tidy_dates(["15-03-26", "31-12-25"])
        self.assertEqual(out, ["2026-03-15", "2025-12-31"])
        self.assertEqual(si.tidy_dates(["5 Mar 2026", "12-May-26", "May 12, 2026"])[0],
                         ["2026-03-05", "2026-05-12", "2026-05-12"])

    def test_a_title_line_of_a_few_cells_is_not_the_header(self):
        headers, rows, first = si.split_header([["Colony", "March 2026", ""] + [""] * 5,
                                                ["Ear tag", "Sex", "DOB", "Strain", "Cage", "Room", "Owner", "Notes"],
                                                ["1", "F", "2026-01-01", "Cre", "10", "B1", "sam", ""]])
        self.assertEqual((headers[0], first), ("Ear tag", 3))
        self.assertEqual(si.tidy_dates(["1.5e-07"])[0], [""])   # (numbers are not dates)
        self.assertEqual(si._cell(1.5e-07), "1.5e-07")

    def test_the_export_s_formula_guard_is_undone_on_import(self):
        rows = si.read_workbook("mice.csv", "Genotype,Note\n'+/+,'=1+1\n'-/-,it's fine\n".encode())["Sheet 1"]
        self.assertEqual(rows[1:], [["+/+", "=1+1"], ["-/-", "it's fine"]])
        headers, _rows, _first = si.split_header([["Colony", "March", "2026", "", ""],
                                                  ["Ear tag", "Sex", "DOB", "Strain", "Cage"],
                                                  ["1", "F", "2026-01-01", "Cre", "10"]])
        self.assertEqual(headers[0], "Ear tag")                  # a title of 3 of 5 cells

    def test_excel_date_numbers_and_iso(self):
        out, _ = si.tidy_dates(["46095", "2026-03-14 00:00", "not a date"])
        self.assertEqual(out[:2], ["2026-03-14", "2026-03-14"])
        self.assertEqual(out[2], "")

    def test_a_lab_s_own_mouse_statuses_win(self):
        from app.app import app
        from app.db import SessionLocal
        from tests.base import execute
        execute("insert into dropdown_options (field_name, option_value, created_at) values ('status', 'Stock', '2026-01-01')")
        with app.test_request_context(), SessionLocal() as s:
            choices = si._mouse_statuses(s)
        self.assertEqual(si.tidy_choice("stock", choices), ("Stock", True))
        self.assertEqual(si.tidy_choice("Sacrificed", choices), ("sac", True))
        execute("delete from dropdown_options where field_name='status' and option_value='Stock'")

    def test_a_total_line_is_not_a_record(self):
        for line in (["TOTAL", "", "7 mice"], ["", "Grand total:", "12"], ["Totals"], ["sum", "3"]):
            self.assertTrue(si.is_total(line), line)
        for line in (["Total RNA", "liver"], ["1001", "total"], [""]):
            self.assertFalse(si.is_total(line), line)

    def test_sexes_and_choices(self):
        self.assertEqual(si.tidy_choice("Male", si.SEXES), ("M", True))
        self.assertEqual(si.tidy_choice("♀", si.SEXES)[0], "F")
        self.assertEqual(si.tidy_choice("Sacrificed", si.MOUSE_STATUSES), ("sac", True))
        self.assertEqual(si.tidy_choice("weird", si.MOUSE_STATUSES), ("weird", False))


class Reading(AppTestCase):
    def test_xlsx_with_a_title_row_and_real_dates(self):
        data = xlsx([["Our plasmids"], ["Name", "Date", "N"], ["pA", datetime(2026, 3, 14), 3.0]])
        sheets = si.read_workbook("x.xlsx", data)
        headers, rows, first = si.split_header(sheets["Sheet1"])
        self.assertEqual(headers, ["Name", "Date", "N"])
        self.assertEqual((rows, first), ([["pA", "2026-03-14", "3"]], 3))

    def test_csv_semicolons_and_old_encodings(self):
        data = "Name;Notes\npA;café\n".encode("cp1252")
        headers, rows, _first = si.split_header(si.read_workbook("x.csv", data)["Sheet 1"])
        self.assertEqual((headers, rows), (["Name", "Notes"], [["pA", "café"]]))

    def test_old_xls_says_what_to_do(self):
        with self.assertRaises(si.ImportProblem) as caught:
            si.read_workbook("old.xls", b"\xd0\xcf\x11\xe0")
        self.assertIn("Save As", str(caught.exception))


class Importing(AppTestCase):
    def upload(self, client, target, filename, data):
        r = client.post(f"/import-sheet/{target}/upload", data={"file": (io.BytesIO(data), filename)},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302, r.get_data(as_text=True)[:500])
        token = r.headers["Location"].rsplit("/", 1)[1]
        return token, self.get_ok(client, f"/import-sheet/file/{token}")

    @staticmethod
    def chosen(html) -> dict[str, str]:
        """{select name: the option selected} on the match page."""
        out = {}
        for name, body in re.findall(r'<select name="((?:map|kind)-\d+)"[^>]*>(.*?)</select>', html, re.S):
            picked = re.search(r'<option value="([^"]*)" selected', body)
            out[name] = picked.group(1) if picked else ""
        return out

    def test_every_database_has_the_button(self):
        self.assertIn("/import-sheet/mice", self.get_ok(self.a, "/colony?view=mice"))
        self.assertIn("/import-sheet/plasmids", self.get_ok(self.a, "/plasmids"))
        self.assertIn("/import-sheet/fish", self.get_ok(self.a, "/zebrafish?view=fish"))
        self.assertIn("/import-sheet/inventory:reagents", self.get_ok(self.a, "/inventory/reagents"))
        key = self.make_stock_module(self.a)
        self.assertIn(f"/import-sheet/stocks:{key}", self.get_ok(self.a, f"/stocks/{key}"))
        key = self.make_organism_module(self.a)
        self.assertIn(f"/import-sheet/organisms:{key}", self.get_ok(self.a, f"/organisms/{key}"))

    def test_mice_from_excel(self):
        tag = uniq("TG")
        data = xlsx([["Ear tag", "Sex", "DOB", "Strain", "Cage #", "Room", "Status", "Cage colour"],
                     ["", "Male", "14/03/2026", tag, "", "B12", "Breeding", "blue"],
                     ["", "f", "03/04/2026", tag, "", "B12", "sacrificed", "red"]])
        token, html = self.upload(self.a, "mice", "colony.xlsx", data)
        chosen = self.chosen(html)
        self.assertEqual(chosen["map-1"], "gender")
        self.assertEqual(chosen["map-2"], "date_of_birth")
        self.assertEqual(chosen["map-3"], "genotype")
        self.assertEqual(chosen["map-5"], "cage_location")       # B12 is a room here, not a box position
        self.assertEqual(chosen["map-6"], "status")
        self.assertEqual(chosen["map-7"], "_notes")               # the colony's columns are fixed
        self.assertIn("isn't in your sheet", html)               # no owner column: every row gets one
        form = {**chosen, "sheet": "Sheet1", "fill-owner": "me"}
        before = count("mice")
        preview = self.a.post(f"/import-sheet/file/{token}/preview", data=form).get_data(as_text=True)
        self.assertIn("2 of 2 rows", preview)
        self.assertEqual(count("mice"), before)                  # a preview writes nothing
        r = self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        self.assertIn("Imported 2 mice", flash_text(r))
        self.assertEqual(count("mice"), before + 2)
        got = row("select gender, status, owner, note from mice order by id desc limit 1")
        self.assertEqual(got[:3], ("F", "sac", self.admin))
        self.assertIn("Cage colour: red", got[3])
        self.assertIn("Room: B12", got[3])                       # no cage to hold the room
        self.assertEqual(last_batch()[2], "create")

    def test_an_ear_tag_column_goes_to_the_ear_tag_not_the_mouse_number(self):
        """It used to be read as the Mouse ID, for want of anywhere else to
        put it; mice have a column for what is written on them now."""
        tag = uniq("TG")
        data = xlsx([["Ear tag", "Sex", "Strain"], ["RF", "M", tag], ["142", "F", tag]])
        token, html = self.upload(self.a, "mice", "colony.xlsx", data)
        form = {**self.chosen(html), "sheet": "Sheet1", "fill-owner": "me"}
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        self.assertEqual(sorted(r[0] for r in rows("select ear_tag from mice where transgene_1=?", tag)),
                         ["142", "RF"])

    def test_a_tail_tattoo_column_goes_to_the_custom_tag(self):
        """However the lab marks its mice, the column goes to Custom tag."""
        for header in ("Tail tattoo", "Custom tag"):
            tag = uniq("TG")
            data = xlsx([[header, "Sex", "Strain"], ["T-12", "M", tag]])
            token, html = self.upload(self.a, "mice", "colony.xlsx", data)
            form = {**self.chosen(html), "sheet": "Sheet1", "fill-owner": "me"}
            self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
            self.assertEqual(one("select ear_tag from mice where transgene_1=?", tag), "T-12", header)

    def test_a_taken_mouse_id_and_a_total_line_are_named(self):
        n = (one("select max(mouse_id) from mice") or 0) + 1000
        tag = uniq("TG")
        data = xlsx([["Mouse ID", "Sex", "Strain"], [str(n), "M", tag], [str(n), "F", tag], ["TOTAL", "", "2 mice"]])
        token, html = self.upload(self.a, "mice", "colony.xlsx", data)
        form = {**self.chosen(html), "sheet": "Sheet1", "fill-owner": "me"}
        preview = self.a.post(f"/import-sheet/file/{token}/preview", data=form).get_data(as_text=True)
        self.assertIn("2 of 2 rows", preview)
        self.assertIn("Row 4 looks like the sheet&#39;s total, so it&#39;s left out.", preview)
        self.assertIn(f"Mouse ID {n} is taken (in the colony or by an earlier row)", preview)
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        self.assertEqual(count("mice", "transgene_1=?", tag), 2)
        second = row("select mouse_id, note from mice where transgene_1=? and gender='F'", tag)
        self.assertNotEqual(second[0], n)
        self.assertIn(f"ID in the spreadsheet: {n}", second[1])

    def import_mice(self, data, filename="colony.xlsx", **fills):
        token, html = self.upload(self.a, "mice", filename, data)
        form = {**self.chosen(html), "sheet": "Sheet1", **fills}
        preview = self.a.post(f"/import-sheet/file/{token}/preview", data=form).get_data(as_text=True)
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        return preview

    def test_a_cage_merged_down_over_its_mice_holds_each_of_them(self):
        tag, first, second = uniq("TG"), uniq("M"), uniq("M")
        data = xlsx([["Cage #", "Sex", "Strain"], [first, "F", tag], [None, "F", tag], [None, "M", tag],
                     [second, "M", tag], [None, "F", tag]], merge=("A2:A4", "A5:A6"))
        self.import_mice(data, **{"fill-owner": "me"})
        cages = [c for (c,) in rows(
            "select c.cage_id from mice m join mouse_cages c on c.id=m.cage_id_fk where m.transgene_1=? "
            "order by m.id", tag)]
        self.assertEqual(cages, [first, first, first, second, second])

    def test_a_new_cage_for_a_rack_place_never_takes_a_number_the_sheet_uses(self):
        from app.db import SessionLocal
        from app.services import reserve_cage_ids
        with SessionLocal() as s:
            upcoming = reserve_cage_ids(s, 1)[0]
        tag = uniq("TG")
        data = xlsx([["Cage #", "Rack", "Position", "Sex", "Strain", "Owner"],
                     ["", "Rack A", "A1", "F", tag, self.member], [upcoming, "", "", "M", tag, self.member]])
        self.import_mice(data)
        placed, typed = [tuple(r) for r in rows(
            "select c.cage_id, c.owner from mice m join mouse_cages c on c.id=m.cage_id_fk "
            "where m.transgene_1=? order by m.id", tag)]
        self.assertNotEqual(placed[0], upcoming)
        self.assertEqual(count("mice", "cage_id_fk=(select id from mouse_cages where cage_id=?)", upcoming), 1)
        self.assertEqual((placed[1], typed[1]), (self.member, self.member))   # the mice's owner's, not the importer's

    def test_dates_follow_the_lab_s_style_and_nothing_is_born_tomorrow(self):
        from app.db import SessionLocal
        from app.inventory_service import set_setting
        def style(value):
            with SessionLocal() as s:
                set_setting(s, "date_style", value)
                s.commit()
        style("day")
        self.addCleanup(style, "month")
        tag = uniq("TG")
        data = (f"Sex,Strain,DOB\nF,{tag},03/04/2026\nM,{tag},12-May-26\nF,{tag},05/06/2099\n").encode()
        preview = self.import_mice(data, "colony.csv", **{"fill-owner": "me"})
        self.assertIn("read day first", preview)
        got = [r for r in rows(
            "select l.date_of_birth, m.note from mice m left join litters l on l.id=m.litter_id_fk "
            "where m.transgene_1=? order by m.id", tag)]
        self.assertEqual([str(d) if d else None for d, _n in got], ["2026-04-03", "2026-05-12", None])
        self.assertIn("2099-06-05", got[2][1])

    def test_one_litter_keeps_one_date_of_birth(self):
        tag, litter = uniq("TG"), uniq("L")
        data = (f"Sex,Strain,Litter,DOB\nF,{tag},{litter},2026-03-01\nM,{tag},{litter},2026-03-09\n").encode()
        preview = self.import_mice(data, "colony.csv", **{"fill-owner": "me"})
        self.assertIn("in an earlier row", preview)
        self.assertEqual(str(one("select date_of_birth from litters where litter_id=?", litter)), "2026-03-01")
        self.assertIn("2026-03-09", one("select note from mice where transgene_1=? and gender='M'", tag))

    def test_plasmids_with_boxes_positions_and_a_location_column(self):
        box = uniq("Box ")
        name = uniq("pImp")
        data = (f"Plasmid #,Construct,Vector,Antibiotic,Box,Location,Freezer,Made by\n"
                f",{name}-1,pUC19,Amp,{box},A1,-80 top,{self.member}\n"
                f",{name}-2,pUC19,Kan,{box},A1,-80 top,nobody-here\n").encode()
        token, html = self.upload(self.a, "plasmids", "plasmids.csv", data)
        chosen = self.chosen(html)
        self.assertEqual(chosen["map-5"], "position")
        self.assertEqual(chosen["map-6"], "location")
        self.assertEqual(chosen["map-7"], "owner")
        form = {**chosen, "sheet": "Sheet 1"}
        preview = self.a.post(f"/import-sheet/file/{token}/preview", data=form).get_data(as_text=True)
        self.assertIn("already holds plasmid", preview)          # both want A1
        self.assertIn("isn&#39;t anyone in the lab", preview)
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        first = row("select owner, backbone, resistance, location, box_row, box_col from plasmids where name=?",
                    f"{name}-1")
        self.assertEqual(first, (self.member, "pUC19", "Amp", "-80 top", 0, 0))
        second = row("select owner, box_row, notes from plasmids where name=?", f"{name}-2")
        self.assertEqual(second[:2], (self.admin, None))         # unknown owner → you; the cell was taken
        self.assertIn("nobody-here", second[2])

    def fresh_inventory(self, preset="reagents") -> str:
        r = self.a.post("/inventory/new", data={"preset": preset, "label": uniq("Chemicals "), "audience": "lab"})
        return r.headers["Location"].split("?")[0].rstrip("/").rsplit("/", 1)[1]

    def test_an_inventory_gains_the_columns_it_lacks(self):
        key = self.fresh_inventory()
        name = uniq("Reagent ")
        data = xlsx([["Product", "Supplier", "Cat. No.", "Expiry date", "Shelf life", "Batch"],
                     [name, "Sigma", "S-1", "2027-01-31", "flammable", "L7"]])
        token, html = self.upload(self.a, f"inventory:{key}", "reagents.xlsx", data)
        chosen = self.chosen(html)
        self.assertEqual([chosen[f"map-{i}"] for i in range(6)],
                         ["name", "vendor", "catalog_number", "expires_on", "_new", "lot"])
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet1"})
        item = row("select vendor, catalog_number, expires_on, lot, attrs from inventory_items where name=?", name)
        self.assertEqual(item[:2], ("Sigma", "S-1"))
        self.assertEqual(str(item[2])[:10], "2027-01-31")
        self.assertEqual(json.loads(item[4])["shelf_life"], "flammable")
        settings = json.loads(one("select settings from inventory_modules where key=?", key))
        self.assertIn("Shelf life", [f["label"] for f in settings["fields"]])

    def test_a_member_cannot_add_columns_to_a_lab_inventory(self):
        name = uniq("Reagent ")
        token, html = self.upload(self.m, "inventory:reagents", "r.csv", f"Name,Odd column\n{name},x\n".encode())
        chosen = self.chosen(html)
        self.assertEqual(chosen["map-1"], "_notes")
        self.post(self.m, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet 1"})
        self.assertIn("Odd column: x", one("select notes from inventory_items where name=?", name))

    def test_fly_vials_into_racks(self):
        key = self.make_stock_module(self.a)
        rack = self.make_stock_rack(self.a, key)
        rack_name = one("select name from stock_racks where id=?", rack)
        gt = uniq("w1118; ")
        data = xlsx([["Stock", "Bloomington #", "Rack", "Slot", "Type"], [gt, "5905", rack_name, "A1", "Stock"]])
        token, html = self.upload(self.a, f"stocks:{key}", "flies.xlsx", data)
        chosen = self.chosen(html)
        self.assertEqual([chosen[f"map-{i}"] for i in range(5)],
                         ["genotype", "stock_number", "rack", "position", "purpose"])
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet1"})
        self.assertEqual(row("select rack_id_fk, rack_row, rack_col from stock_units where genotype=?", gt),
                         (rack, 1, 1))

    def test_fish_make_their_tanks_and_lines(self):
        tank, line = uniq("T"), uniq("Line ")
        data = f"Tank,Line,Number of fish,Gender,DOF\n{tank},{line},12,mix,2026-01-05\n".encode()
        token, html = self.upload(self.a, "fish", "fish.csv", data)
        chosen = self.chosen(html)
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet 1"})
        got = row("select f.count, f.sex, l.name from fish f join tanks t on t.id=f.tank_id_fk "
                  "left join fish_lines l on l.id=f.line_id_fk where t.tank_id=?", tank)
        self.assertEqual(got, (12, "mixed", line))

    def test_an_organism_database_gains_columns_and_housing(self):
        key = self.make_organism_module(self.a)
        data = xlsx([["Newt ID", "Sex", "Tank", "Weight (g)"], ["N-1", "Male", "T-9", "4.5"]])
        token, html = self.upload(self.a, f"organisms:{key}", "newts.xlsx", data)
        chosen = self.chosen(html)
        self.assertEqual(chosen["map-2"], "housing")
        self.assertEqual((chosen["map-3"], chosen["kind-3"]), ("_new", "number"))
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet1"})
        mid = self.organism_module_id(key)
        got = row("select o.sex, h.code, o.attrs from organisms o join organism_housing h on h.id=o.housing_id_fk "
                  "where o.module_id_fk=? and o.code='N-1'", mid)
        self.assertEqual(got[:2], ("male", "T-9"))
        self.assertEqual(float(json.loads(got[2])["weight_g"]), 4.5)

    def test_plasmids_can_all_come_in_as_lab_common(self):
        name = uniq("pCommon")
        token, html = self.upload(self.a, "plasmids", "p.csv", f"Name\n{name}\n".encode())
        self.assertIn('name="fill-is_shared"', html)
        self.post(self.a, f"/import-sheet/file/{token}/run",
                  data={**self.chosen(html), "sheet": "Sheet 1", "fill-is_shared": "1"})
        self.assertEqual(one("select is_shared from plasmids where name=?", name), True)

    def test_an_inventory_sheet_says_which_rows_are_lab_common(self):
        key = self.fresh_inventory()
        mine, common = uniq("Tris "), uniq("PBS ")
        data = xlsx([["Name", "Lab common"], [mine, "no"], [common, "yes"]])
        token, html = self.upload(self.a, f"inventory:{key}", "r.xlsx", data)
        chosen = self.chosen(html)
        self.assertEqual(chosen["map-1"], "is_shared")
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**chosen, "sheet": "Sheet1"})
        self.assertEqual([one("select is_shared from inventory_items where name=?", n) for n in (mine, common)],
                         [False, True])

    def test_problems_name_the_sheet_s_own_row(self):
        data = xlsx([["Plasmids of the lab"], ["Name", "Vector"], ["pOk", "pUC19"], [], ["", "no name"]])
        token, html = self.upload(self.a, "plasmids", "p.xlsx", data)
        self.assertIn("2 rows", html)
        preview = self.a.post(f"/import-sheet/file/{token}/preview",
                              data={**self.chosen(html), "sheet": "Sheet1"}).get_data(as_text=True)
        self.assertIn("Row 5: It has no name.", preview)

    def test_missing_must_haves_stop_the_import(self):
        token, html = self.upload(self.a, "plasmids", "p.csv", b"Vector,Notes\npUC19,x\n")
        r = self.post(self.a, f"/import-sheet/file/{token}/preview",
                      data={"map-0": "backbone", "map-1": "_notes", "sheet": "Sheet 1"})
        self.assertIn("Name", flash_text(r))

    def test_an_import_can_be_undone(self):
        name = uniq("pUndo")
        token, html = self.upload(self.a, "plasmids", "p.csv", f"Name\n{name}\n".encode())
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**self.chosen(html), "sheet": "Sheet 1"})
        self.assertEqual(count("plasmids", "name=?", name), 1)
        self.assertIn("upload is finished", flash_text(self.a.get(f"/import-sheet/file/{token}", follow_redirects=True)))
        batch_id = last_batch()[0]
        self.post(self.a, f"/batches/{batch_id}/undo")
        self.assertEqual(count("plasmids", "name=?", name), 0)

    def test_an_upload_is_only_its_owners(self):
        token, _html = self.upload(self.a, "plasmids", "p.csv", b"Name\npX\n")
        self.assertEqual(self.m.get(f"/import-sheet/file/{token}").status_code, 404)

    # -- sheets as labs keep them (benchmarks/import_sheets.py found these)

    def fields(self, target_key):
        from app.app import app
        from app.db import SessionLocal
        with app.test_request_context(), SessionLocal() as s:
            return si.target_for(s, target_key).fields

    @staticmethod
    def ticked(html) -> dict[str, str]:
        """The checkboxes the match page ticks, as a form sends them."""
        return {name: "1" for name in re.findall(r'<input type="checkbox" name="(down-\d+)" value="1" checked', html)}

    def test_a_cage_written_once_per_cage_fills_the_rows_below(self):
        """A sheet filled in by eye writes the cage on its first mouse only;
        the blanks below are that cage, as merged cells are."""
        tag = uniq("TG")
        data = xlsx([["Cage", "Sex", "Strain"], ["C1" + tag, "M", tag], ["", "F", tag], ["", "F", tag],
                     ["C2" + tag, "M", tag], ["", "M", tag]])
        token, html = self.upload(self.a, "mice", "colony.xlsx", data)
        self.assertIn("Its 3 blank cells take the value above them", html)
        form = {**self.chosen(html), **self.ticked(html), "sheet": "Sheet1", "fill-owner": "me", "down-seen": "1"}
        self.assertEqual(form.get("down-0"), "1")
        preview = self.a.post(f"/import-sheet/file/{token}/preview", data=form).get_data(as_text=True)
        self.assertIn("Cage: 3 blank cells took the value above them.", preview)
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        cages = [r[0] for r in rows("select c.cage_id from mice m join mouse_cages c on c.id = m.cage_id_fk "
                                    "where m.transgene_1=? order by m.id", tag)]
        self.assertEqual(cages, ["C1" + tag] * 3 + ["C2" + tag] * 2)

    def test_unticked_the_blanks_stay_blank(self):
        tag = uniq("TG")
        data = xlsx([["Cage", "Strain"], ["C1" + tag, tag], ["", tag], ["", tag]])
        token, html = self.upload(self.a, "mice", "colony.xlsx", data)
        form = {**self.chosen(html), "sheet": "Sheet1", "fill-owner": "me", "down-seen": "1"}   # no down-0
        self.post(self.a, f"/import-sheet/file/{token}/run", data=form)
        self.assertEqual(count("mice", "transgene_1=? and cage_id_fk is null", tag), 2)
        # Back from the preview ("Change the matches"), it stays as the person left it.
        token, html = self.upload(self.a, "mice", "c.xlsx", data)
        self.assertIn('name="down-0" value="1" checked', html)
        again = self.get_ok(self.a, f"/import-sheet/file/{token}?down-seen=1&map-0=cage_id")
        self.assertNotIn('name="down-0" value="1" checked', again)

    def test_only_a_column_that_looks_filled_down_is_offered(self):
        self.assertEqual(si.filled_down(["101", "", "", "102", ""]), 3)
        self.assertEqual(si.filled_down(["101", "101", "", "102"]), 0)    # written on every row: a blank is no cage
        self.assertEqual(si.filled_down(["101", "", "", "101", "", "102", ""]), 4)   # again after a blank line
        self.assertEqual(si.filled_down(["", "101", ""]), 0)              # nothing above the first blank
        self.assertEqual(si.filled_down(["101", "102", "103"]), 0)
        self.assertEqual(si.fill_down(["a", "", "b", " ", ""]), (["a", "a", "b", "b", "b"], 3))
        tag = uniq("TG")
        html = self.upload(self.a, "mice", "c.xlsx", xlsx([["Notes", "Strain"], ["x", tag], ["", tag]]))[1]
        self.assertNotIn("take the value above", html)                     # notes don't group mice

    def test_chinese_headers_are_matched(self):
        mice = si.auto_match(["小鼠编号", "性别", "出生日期", "笼号", "基因型", "状态", "负责人", "备注"],
                             [["1"], ["公"], ["2026-03-01"], ["101"], ["Ai14/+"], ["种鼠"], ["me"], ["x"]],
                             self.fields("mice"))
        self.assertEqual([k for k, _s, _w in (mice[i] for i in range(8))],
                         ["mouse_id", "gender", "date_of_birth", "cage_id", "genotype", "status", "owner", "note"])
        plasmids = si.auto_match(["质粒名称", "载体", "插入片段", "抗性", "盒子", "位置", "浓度"],
                                 [["pA"], ["pUC19"], ["EGFP"], ["Amp"], ["Box 1"], ["A1"], ["120"]],
                                 self.fields("plasmids"))
        self.assertEqual([k for k, _s, _w in (plasmids[i] for i in range(7))],
                         ["name", "backbone", "insert_seq", "resistance", "box", "position", "concentration"])
        self.assertEqual(si.tidy_choice("公", si.SEXES), ("M", True))
        self.assertEqual(si.tidy_choice("雌", si.SEXES), ("F", True))
        self.assertEqual(si.tidy_choice("处死", si.MOUSE_STATUSES), ("sac", True))

    def test_a_chinese_sheet_of_mice_comes_in(self):
        tag = uniq("TG")
        data = xlsx([["性别", "基因型", "出生日期", "状态"], ["母", tag, "2026-03-01", "种鼠"]])
        token, html = self.upload(self.a, "mice", "小鼠.xlsx", data)
        self.post(self.a, f"/import-sheet/file/{token}/run",
                  data={**self.chosen(html), "sheet": "Sheet1", "fill-owner": "me"})
        self.assertEqual(row("select gender, status from mice where transgene_1=?", tag), ("F", "breeder"))

    def test_a_concentration_may_carry_its_unit(self):
        names = [uniq("pConc") for _ in range(3)]
        data = f"Name,Concentration\n{names[0]},152.6 ng/µl\n{names[1]},98 ng/uL\n{names[2]},1.2 µg/µl\n"
        token, html = self.upload(self.a, "plasmids", "p.csv", data.encode())
        self.post(self.a, f"/import-sheet/file/{token}/run", data={**self.chosen(html), "sheet": "Sheet 1"})
        got = [row("select concentration, notes from plasmids where name=?", n) for n in names]
        self.assertEqual([g[0] for g in got[:2]], ["152.6", "98"])
        self.assertEqual(got[2][0], "")                     # another unit isn't guessed at
        self.assertIn("1.2 µg/µl", got[2][1])               # it's kept in the notes

    def test_a_csv_from_a_chinese_windows_reads_as_chinese(self):
        """Excel there saves CSV in GBK; a Western one in cp1252 stays as it was."""
        tag = uniq("TG")
        data = f"性别,基因型,笼号\r\n公,{tag},101\r\n".encode("gbk")
        headers, rows, _first = si.split_header(si.read_workbook("小鼠.csv", data)["Sheet 1"])
        self.assertEqual((headers, rows), (["性别", "基因型", "笼号"], [["公", tag, "101"]]))
        headers, rows, _first = si.split_header(si.read_workbook("x.csv", "Name,Owner\r\npA,Díaz café\r\n".encode("cp1252"))["Sheet 1"])
        self.assertEqual(rows, [["pA", "Díaz café"]])
