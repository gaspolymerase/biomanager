"""Planned experiments, the plan kept at Start, and the changes from it
(app/experiment_plan.py): an assistant designs one (experiment_new), nothing
of it is due until a person presses Start, and afterwards a change to the
main readout or its day, the groups, an animal left out of the analysis, or
the test asks for a reason and is kept with it."""
from __future__ import annotations

import io
import json
import re
from datetime import date, timedelta

from tests.base import *  # noqa: F401,F403
from tests.base import AppTestCase, count, execute, location, one, rows, uniq

from app import actions
from app.app import app

TODAY = date.today()


def number(mouse_row: int) -> int:
    return one("select mouse_id from mice where id=?", mouse_row)


class Planned(AppTestCase):
    """A colony experiment designed by an assistant: two groups of mice, a
    regimen, the main readout on day 3 and a Welch test."""

    def setUp(self):
        super().setUp()
        self.mice = [self.make_mouse(self.m, self.member) for _ in range(4)]

    def run_actions(self, *changes, apply=True, who=None):
        return actions.run(app, who or self.member, list(changes), apply=apply, description=uniq("P"), source="Claude")

    def design(self, **extra):
        name = uniq("Design ")
        fields = {"name": name, "description": "Does HDM raise the weight loss?",
                  "groups": [{"name": "HDM", "animals": [number(m) for m in self.mice[:2]]},
                             {"name": "vehicle", "animals": [number(m) for m in self.mice[2:]]}],
                  "steps": [{"days": "1-2", "kind": "challenge", "agent": "HDM", "dose": "25 µg", "route": "intranasal",
                             "group": "HDM"},
                            {"days": "3", "kind": "reading"}],
                  "readout": "body_weight", "readout_day": 3, "planned_test": "welch",
                  "start_date": (TODAY - timedelta(days=1)).isoformat(), "notes": "Endpoint: day 3 weight", **extra}
        out = self.run_actions({"action": "experiment_new", "fields": fields})
        self.assertTrue(out.ok, [c.as_dict() for c in out.changes])
        return one("select id from experiments where name=?", name)

    def data(self, exp, client=None):
        return (client or self.m).get(f"/experiments/{exp}/data.json").get_json()

    def start(self, exp, client=None):
        return (client or self.m).post(f"/experiments/{exp}/start")

    def started(self, **extra):
        exp = self.design(**extra)
        self.assertEqual(self.start(exp).status_code, 302)
        return exp

    def deviations(self, exp):
        return count("experiment_deviations", "experiment_id_fk=?", exp)


class Designing(Planned):
    def test_a_preview_keeps_nothing(self):
        name = uniq("Design ")
        out = self.run_actions({"action": "experiment_new", "fields": {"name": name}}, apply=False)
        self.assertTrue(out.ok, [c.as_dict() for c in out.changes])
        self.assertIn("It starts when you press Start", out.changes[0].summary)
        self.assertEqual(count("experiments", "name=?", name), 0)

    def test_the_design_is_a_planned_experiment(self):
        exp = self.design()
        status, desc, plan, test, day = [one(f"select {c} from experiments where id=?", exp) for c in
                                         ("status", "description", "treatment_plan", "planned_test", "readout_day")]
        self.assertEqual((status, test, day), ("planned", "welch", 3))
        self.assertIn("Designed with Claude.", desc)
        self.assertEqual(plan, "Endpoint: day 3 weight")
        data = self.data(exp)
        self.assertEqual(sorted((s["group"], s["label"]) for s in data["subjects"]),
                         sorted([("HDM", f"#{number(m)}") for m in self.mice[:2]]
                                + [("vehicle", f"#{number(m)}") for m in self.mice[2:]]))
        self.assertEqual([s["agent"] for s in data["steps"]], ["HDM", ""])
        self.assertTrue(data["planned"])
        self.assertFalse(data["fixed"])
        # Its days are planned, not due, even those already past.
        self.assertEqual({r["state"] for r in data["schedule"]}, {"planned"})
        page = self.get_ok(self.m, f"/colony/experiments/{exp}")
        self.assertIn(f'action="/experiments/{exp}/start"', page)
        self.assertIn("nothing of it is due yet", page)

    def test_animals_from_another_database_are_refused(self):
        tank = self.make_tank(self.m)
        out = self.run_actions({"action": "experiment_new", "fields": {
            "name": uniq("D "), "groups": [{"name": "x", "animals": [{"kind": "tank", "id": tank}]}]}}, apply=False)
        self.assertFalse(out.ok)
        self.assertIn("isn't an animal", " ".join(out.changes[0].errors))

    def test_nothing_is_due_before_start(self):
        from app import api
        from app import experiment_steps as xs
        from app.db import SessionLocal
        exp = self.design(start_date=TODAY.isoformat())
        with app.test_request_context("/"):
            with SessionLocal() as s:
                items = xs.calendar_items(s, TODAY - timedelta(days=5), TODAY + timedelta(days=10))
        self.assertFalse([i for i in items if i["raw"]["anchor_id"] == exp])
        api.rate.reset()
        execute("delete from app_settings where key='members_api_tokens'")  # another test may have left it off
        r = self.m.post("/settings/api-tokens", data={"label": uniq("T"), "scope": "read", "expires": "90"})
        tok = re.search(r'id="api-token"[^>]*>(bmt_\w+)<', r.get_data(as_text=True)).group(1)
        due = app.test_client().get("/api/v1/due", headers={"Authorization": f"Bearer {tok}"}).get_json()["data"]
        steps = {s["id"] for s in self.data(exp)["steps"]}
        self.assertFalse([d for d in due if (d["ref"] or {}).get("kind") == "experiment_step" and d["ref"]["id"] in steps])
        # Started, today's day is due.
        self.start(exp)
        due = app.test_client().get("/api/v1/due", headers={"Authorization": f"Bearer {tok}"}).get_json()["data"]
        self.assertTrue([d for d in due if (d["ref"] or {}).get("kind") == "experiment_step" and d["ref"]["id"] in steps])

    def test_no_change_starts_it(self):
        exp = self.design()
        name = one("select name from experiments where id=?", exp)
        out = self.run_actions({"action": "experiment_change", "target": name, "fields": {"status": "active"}},
                               apply=False)
        self.assertFalse(out.ok)
        self.assertIn("only a person starts it", " ".join(out.changes[0].errors))
        # Nor the page's own form, by status.
        r = self.autosave(self.m, f"/experiments/{exp}/update", {"status": "active"})
        self.assertRefused(r)
        self.assertIn("press Start", r.get_json()["error"])
        r = self.autosave(self.m, f"/colony/experiments/{exp}/update", {"status": "done"})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(one("select status from experiments where id=?", exp), "planned")
        # A plan can be set aside, though.
        self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update", {"status": "cancelled"}))

    def test_the_test_and_the_readout_day_change_freely_while_planned(self):
        exp = self.design()
        self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update", {"planned_test": "anova",
                                                                              "readout_day": "5"}))
        self.assertEqual(one("select planned_test from experiments where id=?", exp), "anova")
        self.assertEqual(self.deviations(exp), 0)
        self.assertRefused(self.autosave(self.m, f"/experiments/{exp}/update", {"planned_test": "kruskal"}))


class Starting(Planned):
    def test_start_keeps_the_plan(self):
        exp = self.design(start_date=None)
        self.assertIsNone(one("select start_date from experiments where id=?", exp))
        r = self.start(exp)
        self.assertIn(f"/experiments/{exp}", location(r))
        self.assertEqual(one("select status from experiments where id=?", exp), "active")
        self.assertEqual(str(one("select start_date from experiments where id=?", exp))[:10], TODAY.isoformat())
        plan = json.loads(one("select plan_json from experiments where id=?", exp))
        self.assertEqual([(g["name"], g["n"]) for g in plan["groups"]], [("HDM", 2), ("vehicle", 2)])
        self.assertEqual((plan["readout"]["key"], plan["readout_day"], plan["test"]), ("body_weight", 3, "welch"))
        self.assertEqual([s["agent"] for s in plan["steps"]], ["HDM", ""])
        data = self.data(exp)
        self.assertTrue(data["fixed"])
        self.assertEqual(data["deviations"], [])
        self.assertTrue(data["comparison"])
        self.assertFalse(any(r["changed"] for r in data["comparison"]))
        # Once is enough.
        self.start(exp)
        self.assertEqual(self.deviations(exp), 0)
        self.assertNotIn(f'action="/experiments/{exp}/start"', self.get_ok(self.m, f"/colony/experiments/{exp}"))

    def test_only_its_owner_starts_it(self):
        exp = self.design()
        self.start(exp, client=self.o)
        self.assertEqual(one("select status from experiments where id=?", exp), "planned")


class ChangesFromThePlan(Planned):
    def test_the_readout_asks_why(self):
        exp = self.started()
        r = self.autosave(self.m, f"/experiments/{exp}/update", {"readout": "score"})
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.get_json()["needs_reason"])
        self.assertIn("the main readout", r.get_json()["error"])
        self.assertNotIn("score", one("select readout from experiments where id=?", exp))   # Cancel: unchanged
        self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update",
                                       {"readout": "score", "reason": "The scale broke"}))
        what, before, after, reason = [one(f"select {c} from experiment_deviations where experiment_id_fk=?", exp)
                                       for c in ("what", "before", "after", "reason")]
        self.assertEqual((what, before, after, reason), ("readout", "Body weight", "Clinical score", "The scale broke"))
        data = self.data(exp)
        self.assertEqual(data["deviations"][0]["reason"], "The scale broke")
        self.assertTrue(next(r for r in data["comparison"] if r["what"] == "Main readout")["changed"])
        # Notes, dates and the name never ask.
        self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update", {
            "description": "more words", "end_date": (TODAY + timedelta(days=20)).isoformat(), "name": uniq("N ")}))

    def test_the_readout_day_and_the_test_ask_why(self):
        exp = self.started()
        for field, value in (("readout_day", "7"), ("planned_test", "anova")):
            r = self.autosave(self.m, f"/experiments/{exp}/update", {field: value})
            self.assertEqual(r.status_code, 409, field)
            self.assertTrue(r.get_json()["needs_reason"])
            self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update", {field: value, "reason": f"why {field}"}))
        self.assertEqual(sorted(x[0] for x in rows("select what from experiment_deviations where experiment_id_fk=?", exp)),
                         ["readout_day", "test"])
        self.assertEqual(one("select planned_test from experiments where id=?", exp), "anova")

    def test_the_groups_ask_why(self):
        exp = self.started()
        extra = self.make_mouse(self.m, self.member)
        r = self.m.post(f"/experiments/{exp}/subjects/add", json={"how": "one", "value": str(extra), "group": "HDM"})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(count("experiment_mice", "experiment_id_fk=? and mouse_id_fk=?", exp, extra), 0)
        r = self.m.post(f"/experiments/{exp}/subjects/add", json={"how": "one", "value": str(extra), "group": "HDM",
                                                                  "reason": "One HDM mouse died before day 1"})
        self.assertSaved(r)
        self.assertEqual(one("select after from experiment_deviations where experiment_id_fk=?", exp),
                         "HDM: 3 · vehicle: 2")
        # More changes to the groups straight after join it, without asking again.
        key = f"mouse:{self.mice[2]}"
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/{key}/update", json={"group": "HDM"}))
        self.assertEqual(self.deviations(exp), 1)
        detail = one("select detail from experiment_deviations where experiment_id_fk=?", exp)
        self.assertIn(f"+ #{number(extra)} (HDM)", detail)
        self.assertIn(f"#{number(self.mice[2])}: vehicle → HDM", detail)

    def test_moving_an_animal_and_taking_one_out_ask_why(self):
        exp = self.started()
        key = f"mouse:{self.mice[0]}"
        r = self.m.post(f"/experiments/{exp}/subjects/{key}/update", json={"group": "vehicle"})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(one("select treatment_group from experiment_mice where experiment_id_fk=? and mouse_id_fk=?",
                             exp, self.mice[0]), "HDM")
        r = self.m.post(f"/experiments/{exp}/subjects/{key}/remove", json={})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 4)
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/{key}/remove", json={"reason": "Fought"}))
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 3)

    def test_a_group_renamed_whole(self):
        exp = self.started()
        key = f"mouse:{self.mice[0]}"
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/{key}/update",
                                     json={"group": "HDM 25", "whole_group": 1, "reason": "Name the dose"}))
        groups = {one("select treatment_group from experiment_mice where experiment_id_fk=? and mouse_id_fk=?", exp, m)
                  for m in self.mice[:2]}
        self.assertEqual(groups, {"HDM 25"})

    def test_leaving_an_animal_out_says_why_and_the_tests_honour_it(self):
        exp = self.started()
        key = f"mouse:{self.mice[0]}"
        r = self.m.post(f"/experiments/{exp}/subjects/{key}/leave-out", json={})
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.get_json()["needs_reason"])
        data = self.m.post(f"/experiments/{exp}/subjects/{key}/leave-out", json={"reason": "Tumour ulcerated"}).get_json()
        self.assertTrue(data["ok"])
        self.assertEqual(next(s for s in data["subjects"] if s["key"] == key)["left_out"], "Tumour ulcerated")
        self.assertEqual((one("select what from experiment_deviations where experiment_id_fk=?", exp),
                          one("select reason from experiment_deviations where experiment_id_fk=?", exp)),
                         ("left_out", "Tumour ulcerated"))
        # Its weight stays in the sheet and leaves the test.
        weights = {self.mice[0]: "30", self.mice[1]: "20", self.mice[2]: "21", self.mice[3]: "22"}
        data = self.m.post(f"/experiments/{exp}/readings", json={"on": TODAY.isoformat(), "values": {
            f"mouse:{m}": v for m, v in weights.items()}}).get_json()
        row = next(r for r in data["table"]["rows"] if r["key"] == key)
        self.assertEqual(row["values"], [30.0])
        hdm = next(g for g in data["table"]["stats"]["raw"][0]["groups"] if g["group"] == "HDM")
        self.assertEqual((hdm["n"], hdm["value"]), (1, 20.0))
        # Putting it back is a change from the plan too.
        r = self.m.post(f"/experiments/{exp}/subjects/{key}/leave-out", json={"back": 1})
        self.assertEqual(r.status_code, 409)
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/{key}/leave-out",
                                     json={"back": 1, "reason": "Vet: not ulcerated"}))
        self.assertEqual(self.deviations(exp), 2)

    def test_the_export_has_the_plan_and_its_changes(self):
        from openpyxl import load_workbook
        exp = self.started()
        self.autosave(self.m, f"/experiments/{exp}/update", {"planned_test": "anova", "reason": "A third group"})
        r = self.m.get(f"/experiments/{exp}/export.xlsx")
        book = load_workbook(io.BytesIO(r.data))
        self.assertIn("Plan and deviations", book.sheetnames)
        cells = [c for row in book["Plan and deviations"].iter_rows(values_only=True) for c in row if c]
        self.assertIn("Welch's t-test", cells)
        self.assertIn("A third group", cells)
        self.assertIn("Left out of the analysis", [c.value for c in book["Readout"][1]])

    def test_the_api_has_the_plan_and_the_changes(self):
        from app import api
        exp = self.started()
        self.autosave(self.m, f"/experiments/{exp}/update", {"readout_day": "4", "reason": "Weekend"})
        api.rate.reset()
        execute("delete from app_settings where key='members_api_tokens'")  # another test may have left it off
        r = self.m.post("/settings/api-tokens", data={"label": uniq("T"), "scope": "read", "expires": "90"})
        tok = re.search(r'id="api-token"[^>]*>(bmt_\w+)<', r.get_data(as_text=True)).group(1)
        body = app.test_client().get(f"/api/v1/experiments/{exp}", headers={"Authorization": f"Bearer {tok}"}).get_json()
        self.assertEqual(body["planned_test"], "welch")
        self.assertEqual(body["plan"]["readout_day"], 3)
        self.assertEqual([(d["what"], d["before"], d["after"], d["reason"]) for d in body["deviations"]],
                         [("readout_day", "3", "4", "Weekend")])


class AssistantsGiveReasonsToo(Planned):
    def test_adding_animals_needs_a_reason(self):
        exp = self.started()
        extra = self.make_mouse(self.m, self.member)
        change = {"action": "experiment_add_animals", "target": {"kind": "experiment", "id": exp},
                  "fields": {"mice": [number(extra)], "group": "vehicle"}}
        out = self.run_actions(change, apply=False)
        self.assertFalse(out.ok)
        self.assertIn("Say why in reason", " ".join(out.changes[0].errors))
        change["fields"]["reason"] = "Replacing one that died"
        out = self.run_actions(change)
        self.assertTrue(out.ok, [c.as_dict() for c in out.changes])
        self.assertEqual(one("select reason from experiment_deviations where experiment_id_fk=?", exp),
                         "Replacing one that died")

    def test_the_test_needs_a_reason(self):
        exp = self.started()
        change = {"action": "experiment_change", "target": {"kind": "experiment", "id": exp},
                  "fields": {"planned_test": "anova"}}
        out = self.run_actions(change, apply=False)
        self.assertIn("Say why in reason", " ".join(out.changes[0].errors))
        change["fields"]["reason"] = "Three groups after all"
        self.assertTrue(self.run_actions(change).ok)
        self.assertEqual(self.deviations(exp), 1)

    def test_the_colony_selection_bar_asks_too(self):
        exp = self.started()
        extra = self.make_mouse(self.m, self.member)
        r = self.m.post("/colony/mice/bulk-experiment", data={"experiment_id": str(exp), "selected_ids": str(extra)})
        self.assertEqual(count("experiment_mice", "experiment_id_fk=? and mouse_id_fk=?", exp, extra), 0)
        self.m.post("/colony/mice/bulk-experiment", data={"experiment_id": str(exp), "selected_ids": str(extra),
                                                          "reason": "A spare"})
        self.assertEqual(count("experiment_mice", "experiment_id_fk=? and mouse_id_fk=?", exp, extra), 1)
        page = self.get_ok(self.m, "/colony?view=mice")
        self.assertIn("data-ask-reason=", page)


class StartedBeforeThis(AppTestCase):
    """An experiment with no plan copy works as it always did."""

    def test_nothing_asks(self):
        r = self.m.post("/experiments/in/colony/create", data={"name": uniq("Old "), "start_date": TODAY.isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        mouse = self.make_mouse(self.m, self.member)
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/add", json={"how": "one", "value": str(mouse)}))
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/mouse:{mouse}/update", json={"group": "A"}))
        self.assertSaved(self.autosave(self.m, f"/experiments/{exp}/update", {"readout": "score", "planned_test": "welch",
                                                                              "readout_day": "9"}))
        self.assertSaved(self.m.post(f"/experiments/{exp}/subjects/mouse:{mouse}/remove", json={}))
        self.assertEqual(count("experiment_deviations", "experiment_id_fk=?", exp), 0)
        data = self.m.get(f"/experiments/{exp}/data.json").get_json()
        self.assertEqual((data["plan"], data["fixed"], data["deviations"], data["comparison"]), (None, False, [], []))
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(self.m.get(f"/experiments/{exp}/export.xlsx").data))
        self.assertNotIn("Plan and deviations", book.sheetnames)
        page = self.get_ok(self.m, f"/colony/experiments/{exp}")
        self.assertNotIn("/start\"", page)
        # It can't be planned again, either.
        self.assertRefused(self.autosave(self.m, f"/experiments/{exp}/update", {"status": "planned"}))
