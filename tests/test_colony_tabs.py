"""The mouse colony tabs other than the Mice sheet and the cage sheet:
litters, breeding and weaning, experiments, strains, dropdown presets,
Add many (batch mice), the global search and the admin colony overview.

Permission model under test (app/access.py): you edit what you own or what
is unowned; mice in a shared (breeder) cage are everyone's; a litter is
editable by whoever may edit all of its mice; an experiment by its creator;
a strain by whoever added it; presets by admins only. Admins edit anything.
"""
from __future__ import annotations

import io
import re
import unittest

from tests.base import *  # noqa: F401,F403


def mouse_number(row_id: int) -> int:
    """The lab-facing mouse number (what forms like wean-distribute take)."""
    return one("select mouse_id from mice where id=?", row_id)


def cage_of(mouse_row_id: int):
    return one("select cage_id_fk from mice where id=?", mouse_row_id)


def row_html(html: str, marker: str, end: str = "</tr>") -> str:
    """The slice of a page from `marker` to the next `end`."""
    start = html.index(marker)
    return html[start:html.index(end, start)]


# ====================================================================== litters

class LitterCreateTests(AppTestCase):

    def test_create_with_existing_id_is_refused_and_leaves_the_litter_alone(self):
        code = uniq("L")
        self.make_litter(self.a, code, date_of_birth=days_ago(10), cohort_name="original")
        r = self.post(self.a, "/colony/litters/create",
                      {"litter_id": code, "date_of_birth": days_ago(3), "cohort_name": "stale form"})
        self.assertFlash(r, "already exists", "error")
        self.assertEqual(count("litters", "litter_id=?", code), 1)
        self.assertEqual(rows("select date_of_birth, cohort_name from litters where litter_id=?", code),
                         [(days_ago(10), "original")])

    def test_create_without_id_takes_the_next_L_number(self):
        """L-1, L-2, L-3…: one past the highest L-number, plain, unpadded."""
        top = 900000 + int(one("select coalesce(max(id), 0) from litters"))
        self.make_litter(self.a, f"L-{top}", date_of_birth=days_ago(20))
        r = self.post(self.a, "/colony/litters/create", {"litter_id": "", "date_of_birth": T})
        self.assertFlash(r, f"Created litter L-{top + 1}", "success")
        self.assertEqual(one("select date_of_birth from litters where litter_id=?", f"L-{top + 1}"), T)

    def test_litters_named_another_way_do_not_change_the_numbering(self):
        top = 900000 + int(one("select coalesce(max(id), 0) from litters"))
        self.make_litter(self.a, f"L-{top}", date_of_birth=days_ago(20))
        self.make_litter(self.a, uniq("QA") + "-0041", date_of_birth=days_ago(5))
        self.make_litter(self.a, str(top + 50), date_of_birth=days_ago(5))
        r = self.post(self.a, "/colony/litters/create", {"litter_id": "", "date_of_birth": T})
        self.assertFlash(r, f"Created litter L-{top + 1}", "success")


class LitterUpdateTests(AppTestCase):

    def test_member_cannot_change_dob_of_a_litter_whose_mice_are_someone_elses(self):
        col = self.make_colony(self.o, self.other, dob=days_ago(40))
        r = self.autosave(self.m, f"/colony/litters/{col['litter_id']}/update",
                          {"date_of_birth": days_ago(5), "cohort_name": "hacked"})
        self.assertRefused(r)
        self.assertIn(self.other, r.get_json()["error"])
        self.assertEqual(rows("select date_of_birth, cohort_name from litters where id=?", col["litter_id"]),
                         [(days_ago(40), "")])

    def test_member_may_edit_a_litter_whose_mice_sit_in_a_shared_breeder_cage(self):
        col = self.make_colony(self.o, self.other, purpose="Breeder")
        r = self.autosave(self.m, f"/colony/litters/{col['litter_id']}/update", {"cohort_name": "shared cohort"})
        self.assertSaved(r)
        self.assertEqual(one("select cohort_name from litters where id=?", col["litter_id"]), "shared cohort")

    def test_member_may_edit_a_litter_of_their_own_mice(self):
        col = self.make_colony(self.m, self.member, dob=days_ago(30))
        r = self.autosave(self.m, f"/colony/litters/{col['litter_id']}/update",
                          {"date_of_birth": days_ago(31), "father_info": "M1", "mother_info": "F2"})
        self.assertSaved(r)
        self.assertEqual(rows("select date_of_birth, father_info, mother_info from litters where id=?", col["litter_id"]),
                         [(days_ago(31), "M1", "F2")])

    def test_an_empty_litter_is_editable_by_anyone(self):
        litter = self.make_litter(self.o)
        self.assertSaved(self.autosave(self.m, f"/colony/litters/{litter}/update", {"notes": "by member"}))
        self.assertEqual(one("select notes from litters where id=?", litter), "by member")

    def test_admin_may_edit_any_litter(self):
        col = self.make_colony(self.o, self.other)
        self.assertSaved(self.autosave(self.a, f"/colony/litters/{col['litter_id']}/update", {"cohort_name": "admin"}))
        self.assertEqual(one("select cohort_name from litters where id=?", col["litter_id"]), "admin")

    def test_update_leaves_fields_absent_from_the_form_alone(self):
        litter = self.make_litter(self.m, date_of_birth=days_ago(12), cohort_name="keep", notes="keep too")
        self.assertSaved(self.autosave(self.m, f"/colony/litters/{litter}/update", {"father_info": "7"}))
        self.assertEqual(rows("select date_of_birth, cohort_name, notes, father_info from litters where id=?", litter),
                         [(days_ago(12), "keep", "keep too", "7")])

    def test_total_pups_must_be_a_whole_number(self):
        litter = self.make_litter(self.m)
        self.assertSaved(self.autosave(self.m, f"/colony/litters/{litter}/update", {"total_pups": "7"}))
        r = self.autosave(self.m, f"/colony/litters/{litter}/update", {"total_pups": "seven"})
        self.assertRefused(r)
        self.assertIn("whole number", r.get_json()["error"])
        self.assertEqual(one("select total_pups from litters where id=?", litter), 7)
        self.assertSaved(self.autosave(self.m, f"/colony/litters/{litter}/update", {"total_pups": ""}))
        self.assertEqual(one("select total_pups from litters where id=?", litter), 0)

    # A litter's DOB is the age of every mouse in it: adding a mouse (Add
    # mouse, Add many) can't re-date a litter whose mice you can't all edit.
    def test_member_cannot_change_a_locked_litters_dob_by_adding_a_mouse_to_it(self):
        col = self.make_colony(self.o, self.other, dob=days_ago(40))
        self.m.post("/colony/mice/create", data={
            "cage_id": "", "litter_id": col["litter"], "date_of_birth": days_ago(2),
            "status": "experiment", "gender": "F", "owner": self.member})
        self.assertEqual(one("select date_of_birth from litters where id=?", col["litter_id"]), days_ago(40))

    def test_litters_tab_locks_rows_the_viewer_may_not_edit(self):
        theirs = self.make_colony(self.o, self.other)
        mine = self.make_colony(self.m, self.member)
        html = self.get_ok(self.m, "/colony?view=litters")
        locked = row_html(html, f'data-litter_id="{theirs["litter"]}"')
        open_ = row_html(html, f'data-litter_id="{mine["litter"]}"')
        self.assertIn("is-locked", locked)
        self.assertRegex(locked, r'name="date_of_birth"[^>]*disabled')
        self.assertNotIn("is-locked", open_)
        self.assertNotRegex(open_, r'name="date_of_birth"[^>]*disabled')


class LitterAddExistingMouseTests(AppTestCase):

    def test_member_cannot_put_someone_elses_mouse_into_a_litter(self):
        litter = self.make_litter(self.m)
        theirs = self.make_mouse(self.o, self.other, cage=uniq("C"))
        before = one("select litter_id_fk from mice where id=?", theirs)
        r = self.post(self.m, f"/colony/litters/{litter}/add-existing-mouse", {"mouse_id": mouse_number(theirs)})
        self.assertFlash(r, "was not added", "error")
        self.assertEqual(one("select litter_id_fk from mice where id=?", theirs), before)

    def test_member_adds_their_own_mouse_to_an_editable_litter(self):
        code = uniq("L")
        litter = self.make_litter(self.m, code)
        mine = self.make_mouse(self.m, self.member, cage=uniq("C"))
        r = self.post(self.m, f"/colony/litters/{litter}/add-existing-mouse", {"mouse_id": mouse_number(mine)})
        self.assertFlash(r, f"is now in litter {code}", "success")
        self.assertEqual(one("select litter_id_fk from mice where id=?", mine), litter)

    def test_a_locked_litter_refuses_even_the_members_own_mouse(self):
        col = self.make_colony(self.o, self.other)
        mine = self.make_mouse(self.m, self.member, cage=uniq("C"))
        r = self.post(self.m, f"/colony/litters/{col['litter_id']}/add-existing-mouse", {"mouse_id": mouse_number(mine)})
        self.assertFlash(r, "has mice owned by", "error")
        self.assertIsNone(one("select litter_id_fk from mice where id=?", mine))

    def test_unknown_or_malformed_mouse_number_is_reported(self):
        litter = self.make_litter(self.m)
        self.assertFlash(self.post(self.m, f"/colony/litters/{litter}/add-existing-mouse", {"mouse_id": "abc"}),
                         "Enter the number of an existing mouse", "error")
        missing = one("select coalesce(max(mouse_id), 0) + 1000 from mice")
        self.assertFlash(self.post(self.m, f"/colony/litters/{litter}/add-existing-mouse", {"mouse_id": missing}),
                         "was not found", "error")


# ============================================================ breeding & weaning

class GiveBirthGenotypingTests(AppTestCase):

    def test_give_birth_stamps_today_on_the_users_cage(self):
        cage = self.make_cage(self.m)
        self.m.post(f"/colony/cages/{cage}/give-birth")
        self.assertEqual(one("select date_give_birth from mouse_cages where id=?", cage), T)

    def test_the_birth_date_confirmed_in_the_dialog_is_recorded(self):
        cage = self.make_cage(self.m)
        r = self.post(self.m, f"/colony/cages/{cage}/give-birth", {"date_give_birth": days_ago(2)})
        self.assertEqual(one("select date_give_birth from mouse_cages where id=?", cage), days_ago(2))
        self.assertFlash(r, "Recorded a litter born on", "success")

    def test_a_birth_date_in_the_future_is_refused(self):
        cage = self.make_cage(self.m)
        r = self.post(self.m, f"/colony/cages/{cage}/give-birth",
                      {"date_give_birth": (date.today() + timedelta(days=3)).isoformat()})
        self.assertFlash(r, "can't be in the future", "error")
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_litter_born_asks_for_the_date(self):
        cage = self.make_cage(self.m, purpose="Breeder")
        html = self.get_ok(self.m, "/colony?view=cages&scope=all")
        button = re.search(rf'<button[^>]*data-litter-born\s+data-action="/colony/cages/{cage}/give-birth"[^>]*>(.*?)</button>',
                           html, re.S)
        self.assertIsNotNone(button)
        self.assertIn("Litter born", button.group(1))
        self.assertNotIn("today", button.group(1))
        self.assertIn('id="litter-born-date" name="date_give_birth" required', html)

    def test_member_cannot_record_a_birth_in_someone_elses_cage(self):
        cage = self.make_cage(self.o, purpose="Experiments")
        self.m.post(f"/colony/cages/{cage}/give-birth")
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_a_breeder_cage_in_any_case_is_shared_but_a_breeding_one_is_its_owners(self):
        # Only a breeder cage starts shared (access.STARTS_SHARED_PURPOSES), so
        # a member may record a birth in someone else's, not in their breeding cage.
        for purpose, shared in (("Breeder", True), ("BREEDER", True), ("Breeding", False)):
            with self.subTest(purpose=purpose):
                cage = self.make_cage(self.o, purpose=purpose)
                self.m.post(f"/colony/cages/{cage}/give-birth")
                self.assertEqual(one("select date_give_birth from mouse_cages where id=?", cage), T if shared else None)

    def test_wean_clears_the_birth_date(self):
        cage = self.make_cage(self.m, date_give_birth=days_ago(21))
        self.m.post(f"/colony/cages/{cage}/wean")
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_genotyping_creates_a_litter_of_pups_in_the_cage(self):
        code = uniq("C")
        cage = self.make_cage(self.m, code, date_give_birth=days_ago(8))
        r = self.post(self.m, f"/colony/cages/{cage}/genotyping",
                      {"total_pups": "3", "father_info": "11", "mother_info": "12"})
        self.assertFlash(r, "with 3 pups", "success")
        litter = row("select id, date_of_birth, total_pups, father_info, mother_info, cohort_name "
                     "from litters order by id desc limit 1")
        self.assertEqual(litter[1:], (days_ago(8), 3, "11", "12", f"Cage {code}"))
        pups = rows("select status, owner, cage_id_fk from mice where litter_id_fk=?", litter[0])
        self.assertEqual(pups, [("geno", self.member, cage)] * 3)

    def test_position_columns_are_short_and_named_in_full(self):
        self.make_cage(self.m)
        for view, key in (("mice", "cage_position"), ("cages", "position")):
            with self.subTest(view=view):
                html = self.get_ok(self.m, f"/colony?view={view}")
                th = re.search(r'<th[^>]*data-sort-key="' + key + r'"[^>]*>.*?</th>', html, re.S).group(0)
                self.assertIn('data-dt-label="Position"', th)
                self.assertIn('title="Position in the rack', th)
                self.assertIn(">Pos<", th)

    def test_genotyping_dialog_says_what_to_type(self):
        self.make_cage(self.m, purpose="Breeder")
        html = self.get_ok(self.m, "/colony?view=cages")
        dialog = html.split('id="genotyping-modal"', 1)[1].split("</dialog>", 1)[0]
        self.assertIn("Type the mouse numbers of the father and the mother, and how many pups were born.", dialog)

    def test_genotyping_rejects_pups_that_are_not_1_to_40(self):
        cage = self.make_cage(self.m)
        litters, mice = count("litters"), count("mice")
        for raw in ("abc", "0", "41", ""):
            with self.subTest(total_pups=raw):
                r = self.post(self.m, f"/colony/cages/{cage}/genotyping", {"total_pups": raw})
                self.assertFlash(r, "whole number from 1 to 40", "error")
        self.assertEqual((count("litters"), count("mice")), (litters, mice))

    def test_member_cannot_genotype_someone_elses_cage(self):
        cage = self.make_cage(self.o, purpose="Experiments")
        mice = count("mice")
        self.m.post(f"/colony/cages/{cage}/genotyping", data={"total_pups": "2"})
        self.assertEqual(count("mice"), mice)


class WeanDueTests(AppTestCase):
    """Weaning is P21 (services.WEAN_OFFSET_DAYS), from the cage's birth
    date or, failing that, from its pups' litter."""

    def wean_due(self, cage, **fields):
        r = self.autosave(self.m, f"/colony/cages/{cage}/update", fields)
        self.assertSaved(r)
        values = r.get_json()["row"]["values"]
        return values["wean_due"], values["wean_state"]

    def test_wean_due_is_21_days_after_the_birth_date_soon_then_overdue(self):
        cage = self.make_cage(self.m)
        self.assertEqual(self.wean_due(cage, date_give_birth=days_ago(5)), (days_ahead(16), ""))
        self.assertEqual(self.wean_due(cage, date_give_birth=days_ago(19)), (days_ahead(2), "soon"))
        self.assertEqual(self.wean_due(cage, date_give_birth=days_ago(25)), (days_ago(4), "overdue"))

    def test_wean_due_falls_back_to_the_pups_litter(self):
        col = self.make_colony(self.m, self.member, n_mice=1, dob=days_ago(10))
        self.assertEqual(self.wean_due(col["cage_id"], notes="x"), (days_ahead(11), ""))


class OneWeaningListTests(AppTestCase):
    """Home, the calendar and the cage sheet read one list of what is due
    to be weaned (services.weaning_due), and weaning takes a litter off it."""

    def calendar_weans(self, client):
        r = client.get("/calendar/events.json", query_string={"start": days_ago(30), "end": days_ahead(30)})
        return [i["title"] for i in r.get_json()["items"] if i["id"].endswith("-wean")]

    def test_a_cage_birth_date_is_due_on_home_and_the_calendar(self):
        code = uniq("C")
        self.make_cage(self.m, code, date_give_birth=days_ago(18))
        self.assertIn(f"Cage {code}", self.get_ok(self.m, "/home"))
        self.assertTrue(any(f"Cage {code}" in t for t in self.calendar_weans(self.m)))

    def test_a_litter_without_a_cage_date_is_due_on_the_calendar_too(self):
        col = self.make_colony(self.m, self.member, n_mice=2, dob=days_ago(18))
        titles = self.calendar_weans(self.m)
        self.assertIn(f"Wean - Litter {col['litter']} · cage {col['cage']}", titles)
        self.assertIn(f"Litter {col['litter']}", self.get_ok(self.m, "/home"))

    def test_a_litter_in_a_cage_with_its_date_counts_once(self):
        col = self.make_colony(self.m, self.member, n_mice=2, dob=days_ago(18), date_give_birth=days_ago(18))
        titles = [t for t in self.calendar_weans(self.m) if col["cage"] in t]
        self.assertEqual(titles, [f"Wean - Litter {col['litter']} · cage {col['cage']}"])

    def test_weaning_takes_the_litter_off_every_list(self):
        col = self.make_colony(self.m, self.member, n_mice=2, dob=days_ago(22))
        self.m.post(f"/colony/cages/{col['cage_id']}/wean")
        self.assertEqual(one("select weaned_on from litters where id=?", col["litter_id"]), T)
        self.assertNotIn(f"Litter {col['litter']}", self.get_ok(self.m, "/home"))
        self.assertFalse([t for t in self.calendar_weans(self.m) if col["litter"] in t])
        r = self.autosave(self.m, f"/colony/cages/{col['cage_id']}/update", {"notes": "x"})
        self.assertEqual(r.get_json()["row"]["values"]["wean_due"], "")

    def test_distributing_marks_the_litter_weaned_before_the_pups_leave(self):
        col = self.make_colony(self.m, self.member, n_mice=1, dob=days_ago(21), date_give_birth=days_ago(21))
        self.post(self.m, f"/colony/cages/{col['cage_id']}/wean-distribute", {
            "mouse_ids[]": [str(mouse_number(col["mice"][0]))], "gender[]": ["F"], "cage_id[]": [""], "card_id[]": [""]})
        self.assertEqual(one("select weaned_on from litters where id=?", col["litter_id"]), T)
        # Nor is the new cage due to wean, though the pup's litter is young.
        r = self.autosave(self.m, f"/colony/cages/{cage_of(col['mice'][0])}/update", {"notes": "x"})
        self.assertEqual(r.get_json()["row"]["values"]["wean_due"], "")

    def test_weaning_and_distributing_can_be_undone(self):
        col = self.make_colony(self.m, self.member, n_mice=1, dob=days_ago(21), date_give_birth=days_ago(21))
        before = cage_of(col["mice"][0])
        self.post(self.m, f"/colony/cages/{col['cage_id']}/wean-distribute", {
            "mouse_ids[]": [str(mouse_number(col["mice"][0]))], "gender[]": ["F"], "cage_id[]": [""], "card_id[]": [""]})
        self.assertNotEqual(cage_of(col["mice"][0]), before)
        batch = one("select id from batches where description like ? order by id desc", f"wean cage {col['cage']}%")
        self.post(self.m, f"/batches/{batch}/undo")
        self.assertEqual(cage_of(col["mice"][0]), before)
        self.assertEqual(one("select date_give_birth from mouse_cages where id=?", col["cage_id"]), days_ago(21))
        self.assertIsNone(one("select weaned_on from litters where id=?", col["litter_id"]))

    def test_young_pups_are_weaned_only_once_confirmed(self):
        cage = self.make_cage(self.m, date_give_birth=days_ago(12))
        r = self.post(self.m, f"/colony/cages/{cage}/wean", {})
        self.assertFlash(r, "are 12 days old", "error")
        self.assertEqual(one("select date_give_birth from mouse_cages where id=?", cage), days_ago(12))
        self.post(self.m, f"/colony/cages/{cage}/wean", {"early": "1"})
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_the_wean_button_carries_the_pups_by_sex(self):
        col = self.make_colony(self.m, self.member, n_mice=2, dob=days_ago(20))
        html = self.get_ok(self.m, "/colony?view=cages")
        button = re.search(r'<button[^>]*data-wean\s[^>]*data-target-label="' + col["cage"] + r'"[^>]*>', html).group(0)
        self.assertIn('data-age="20"', button)
        ids = ", ".join(str(mouse_number(m)) for m in col["mice"])
        self.assertIn(f"&#34;F&#34;: &#34;{ids}&#34;", button)


class FutureBirthTests(AppTestCase):
    """Nothing is born tomorrow: a future birth date is refused."""

    def test_a_mouse_born_tomorrow_is_refused(self):
        before = count("mice")
        r = self.post(self.m, "/colony/mice/create", {"date_of_birth": days_ahead(1), "owner": self.member})
        self.assertFlash(r, "can't be in the future", "error")
        self.assertEqual(count("mice"), before)

    def test_editing_a_date_of_birth_into_the_future_is_refused(self):
        mouse = self.make_mouse(self.m, self.member, date_of_birth=days_ago(30))
        r = self.autosave(self.m, f"/colony/mice/{mouse}/update", {"date_of_birth": days_ahead(3)})
        self.assertFalse(r.get_json()["ok"])
        self.assertEqual(one("select l.date_of_birth from mice m join litters l on l.id=m.litter_id_fk where m.id=?",
                             mouse), days_ago(30))

    def test_a_litter_or_a_cage_born_in_the_future_is_refused(self):
        litter = self.make_litter(self.m, date_of_birth=days_ago(3))
        self.autosave(self.m, f"/colony/litters/{litter}/update", {"date_of_birth": days_ahead(2)})
        self.assertEqual(one("select date_of_birth from litters where id=?", litter), days_ago(3))
        cage = self.make_cage(self.m)
        self.autosave(self.m, f"/colony/cages/{cage}/update", {"date_give_birth": days_ahead(2)})
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_add_many_shows_the_rows_again_instead_of_creating(self):
        tg = uniq("Tg")
        r = self.m.post("/colony/mice/batch/create", data={
            "rows-0-gender": "F", "rows-0-transgene_1": tg, "rows-0-date_of_birth": days_ahead(5),
            "rows-1-gender": "M", "rows-1-transgene_1": tg, "rows-1-date_of_birth": days_ago(5)})
        html = r.get_data(as_text=True)
        self.assertIn("Row 1: the date of birth is in the future", html)
        self.assertIn(f'value="{tg}"', html)
        self.assertEqual(count("mice", "transgene_1=?", tg), 0)


class BreedersTabTests(AppTestCase):

    def test_breeders_tab_lists_breeder_and_breeding_cages_only(self):
        breeder, breeding, other = uniq("C"), uniq("C"), uniq("C")
        self.make_cage(self.a, breeder, purpose="breeder")
        self.make_cage(self.a, breeding, purpose="Breeding")
        self.make_cage(self.a, other, purpose="Experiments")
        html = self.get_ok(self.m, "/colony?view=breeders&scope=all")
        self.assertIn(f">{breeder}</a></strong>", html)
        self.assertIn(f">{breeding}</a></strong>", html)
        self.assertNotIn(f">{other}</a></strong>", html)


class WhereThingsAreTests(AppTestCase):
    """Cards, the export and the Breeders tab say which rack and position
    a cage is in, and how many of each sex it holds."""

    def placed_cage(self, purpose="breeder"):
        rack = uniq("Rack")
        self.a.post("/colony/racks/save", data={"id": "", "name": rack, "rows": "4", "cols": "6"})
        code = uniq("C")
        rack_id = one("select id from mouse_racks where name=?", rack)
        self.make_cage(self.a, code, purpose=purpose, rack_id=str(rack_id), position="B3")
        self.make_mouse(self.a, self.admin, cage=code, gender="F")
        self.make_mouse(self.a, self.admin, cage=code, gender="F")
        self.make_mouse(self.a, self.admin, cage=code, gender="M")
        return rack, code

    def test_one_cage_per_place_even_when_two_arrive_at_once(self):
        rack, code = self.placed_cage()
        rack_id = one("select id from mouse_racks where name=?", rack)
        other = self.make_cage(self.a, uniq("C"))
        # What a second drop at the same moment would write: refused by the database.
        with self.assertRaises(Exception):
            execute("update mouse_cages set rack_id_fk=?, rack_row=2, rack_col=3 where id=?", rack_id, other)
        # A drop onto the taken place is a swap, in steps the rule allows.
        placed = one("select id from mouse_cages where cage_id=?", code)
        execute("update mouse_cages set rack_id_fk=?, rack_row=1, rack_col=1 where id=?", rack_id, other)
        r = self.a.post(f"/colony/cages/{other}/place", data={"rack_id": rack_id, "row": 2, "col": 3})
        self.assertTrue(r.get_json()["ok"])
        self.assertEqual(row("select rack_row, rack_col from mouse_cages where id=?", placed), (1, 1))
        self.assertEqual(row("select rack_row, rack_col from mouse_cages where id=?", other), (2, 3))

    def test_a_cage_card_says_where_it_goes_and_the_sexes(self):
        rack, code = self.placed_cage()
        html = self.get_ok(self.a, f"/labels/cards/cages?scope=all&ids={one('select id from mouse_cages where cage_id=?', code)}")
        self.assertIn(f"{rack} · B3", html)
        self.assertIn("3 · 2♀ 1♂", html)

    def test_the_export_has_rack_and_position(self):
        rack, code = self.placed_cage()
        text = self.a.get("/colony/mice/export?format=csv&scope=all").get_data(as_text=True)
        header, *lines = text.splitlines()
        self.assertIn("Cage_ID,Rack,Position,", header)
        self.assertTrue(any(f",{code},{rack},B3," in line for line in lines))

    def test_exports_are_a_real_workbook_and_never_run_a_formula(self):
        import io
        from openpyxl import load_workbook
        col = self.make_colony(self.a, self.admin, n_mice=1)
        execute("update mice set note=? where id=?", '=HYPERLINK("http://x.test","click")', col["mice"][0])
        text = self.a.get("/colony/mice/export?format=csv&scope=all").get_data(as_text=True)
        self.assertTrue(text.startswith("\ufeff"))                    # Excel opens it as UTF-8
        self.assertIn("'=HYPERLINK", text)
        r = self.a.get("/colony/mice/export?format=excel&scope=all")
        self.assertIn("mice_export.xlsx", r.headers["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(r.get_data())).active
        cells = [c for row in sheet.iter_rows() for c in row if c.value == '=HYPERLINK("http://x.test","click")']
        self.assertEqual([c.data_type for c in cells], ["s"])           # text, not a formula
        token = self.a.post("/import-sheet/mice/upload", data={"file": (io.BytesIO(r.get_data()), "mice_export.xlsx")},
                            content_type="multipart/form-data")
        self.assertEqual(token.status_code, 302)                       # Import from Excel reads it back

    def test_the_breeders_tab_shows_rack_position_and_sexes(self):
        rack, code = self.placed_cage()
        html = self.get_ok(self.a, "/colony?view=breeders&scope=all")
        self.assertIn(f"{rack} · B3", html)
        self.assertIn("3 alive (2♀ 1♂)", html)


class WeanDistributeTests(AppTestCase):

    def distribute(self, client, cage, *groups):
        """groups: (mouse numbers, gender, existing cage code, card label)."""
        data = {"mouse_ids[]": [], "gender[]": [], "cage_id[]": [], "card_id[]": []}
        for ids, gender, cage_code, card in groups:
            data["mouse_ids[]"].append(",".join(str(i) for i in ids))
            data["gender[]"].append(gender)
            data["cage_id[]"].append(cage_code)
            data["card_id[]"].append(card)
        return self.post(client, f"/colony/cages/{cage}/wean-distribute", data)

    def test_pups_go_to_a_new_cage_owned_by_the_user_with_card_and_sex(self):
        col = self.make_colony(self.m, self.member, n_mice=2, date_give_birth=days_ago(21))
        a, b = col["mice"]
        card = uniq("W")
        r = self.distribute(self.m, col["cage_id"], ([mouse_number(a), mouse_number(b)], "M", "", card))
        self.assertFlash(r, "Distributed 2 mice", "success")
        new_cage = cage_of(a)
        self.assertEqual(cage_of(b), new_cage)
        self.assertNotEqual(new_cage, col["cage_id"])
        self.assertEqual(row("select owner, card_id from mouse_cages where id=?", new_cage), (self.member, card))
        self.assertEqual(rows("select gender from mice where id in (?, ?)", a, b), [("M",), ("M",)])
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", col["cage_id"]))

    def test_only_mice_in_the_source_cage_move(self):
        col = self.make_colony(self.m, self.member, n_mice=1)
        elsewhere = self.make_mouse(self.m, self.member, cage=uniq("C"), gender="F")
        home = cage_of(elsewhere)
        r = self.distribute(self.m, col["cage_id"], ([mouse_number(elsewhere)], "M", "", ""))
        self.assertFlash(r, f"mouse {mouse_number(elsewhere)} is not in cage {col['cage']}", "error")
        self.assertEqual(rows("select cage_id_fk, gender from mice where id=?", elsewhere), [(home, "F")])

    def test_mice_the_user_may_not_edit_stay_in_the_source_cage(self):
        col = self.make_colony(self.m, self.member, n_mice=1)
        # An admin parked someone else's mouse in the member's cage.
        theirs = self.make_mouse(self.a, self.other, cage=col["cage"])
        r = self.distribute(self.m, col["cage_id"], ([mouse_number(theirs), mouse_number(col["mice"][0])], "", "", ""))
        self.assertFlash(r, f"mouse {mouse_number(theirs)} is {self.other}", "error")
        self.assertEqual(cage_of(theirs), col["cage_id"])
        self.assertNotEqual(cage_of(col["mice"][0]), col["cage_id"])

    def test_refuses_an_existing_destination_cage_the_user_may_not_edit(self):
        col = self.make_colony(self.m, self.member, n_mice=1)
        target = uniq("C")
        self.make_cage(self.o, target, purpose="Experiments")
        r = self.distribute(self.m, col["cage_id"], ([mouse_number(col["mice"][0])], "", target, ""))
        self.assertFlash(r, f"cage {target} is {self.other}", "error")
        self.assertEqual(cage_of(col["mice"][0]), col["cage_id"])

    def test_moves_into_an_existing_cage_the_user_may_edit(self):
        col = self.make_colony(self.m, self.member, n_mice=1)
        target_code = uniq("C")
        target = self.make_cage(self.m, target_code)
        self.distribute(self.m, col["cage_id"], ([mouse_number(col["mice"][0])], "F", target_code, ""))
        self.assertEqual(cage_of(col["mice"][0]), target)

    def test_a_typed_cage_id_that_does_not_exist_is_created_for_the_user(self):
        col = self.make_colony(self.m, self.member, n_mice=1)
        code = uniq("C")
        self.distribute(self.m, col["cage_id"], ([mouse_number(col["mice"][0])], "", code, ""))
        self.assertEqual(row("select id, owner from mouse_cages where cage_id=?", code),
                         (cage_of(col["mice"][0]), self.member))

    def test_bad_tokens_are_reported_and_the_cage_is_still_weaned(self):
        cage = self.make_cage(self.m, date_give_birth=days_ago(21))
        missing = one("select coalesce(max(mouse_id), 0) + 1000 from mice")
        r = self.distribute(self.m, cage, (["x1", missing], "", "", ""))
        text = flash_text(r)
        self.assertIn("weaned (no mice were moved)", text)
        self.assertIn("“x1” is not a mouse number", text)
        self.assertIn(f"mouse {missing} does not exist", text)
        self.assertIsNone(one("select date_give_birth from mouse_cages where id=?", cage))

    def test_member_cannot_distribute_from_someone_elses_cage(self):
        col = self.make_colony(self.o, self.other, n_mice=1, purpose="Experiments")
        self.distribute(self.m, col["cage_id"], ([mouse_number(col["mice"][0])], "", "", ""))
        self.assertEqual(cage_of(col["mice"][0]), col["cage_id"])


# ================================================================== experiments

class ExperimentCreateTests(AppTestCase):

    def create(self, client, **fields):
        name = fields.pop("name", None) or uniq("Exp ")
        r = client.post("/colony/experiments/create", data={"name": name, **fields})
        exp = one("select max(id) from experiments where name=?", name)
        return r, exp

    def test_create_makes_the_creator_the_owner_and_opens_the_experiment(self):
        r, exp = self.create(self.m, description="d", treatment_plan="drug X", start_date=days_ago(1))
        self.assertEqual(location(r), f"/colony/experiments/{exp}")
        self.assertEqual(row("select owner_username, status, start_date, treatment_plan from experiments where id=?", exp),
                         (self.member, "active", days_ago(1), "drug X"))

    def test_seeding_from_a_cage_skips_mice_the_user_may_not_edit(self):
        col = self.make_colony(self.m, self.member, n_mice=2)
        theirs = self.make_mouse(self.a, self.other, cage=col["cage"])
        name = uniq("Exp ")
        r = self.post(self.m, "/colony/experiments/create", {"name": name, "from_cage_id": col["cage"]})
        exp = one("select id from experiments where name=?", name)
        self.assertFlash(r, "1 skipped", "error")
        members = {m for (m,) in rows("select mouse_id_fk from experiment_mice where experiment_id_fk=?", exp)}
        self.assertEqual(members, set(col["mice"]))
        self.assertNotIn(theirs, members)

    def test_seeding_skips_dead_mice(self):
        col = self.make_colony(self.m, self.member, n_mice=2)
        self.m.post("/colony/mice/bulk-update", data={"field": "status", "value": "sac", "selected_ids": [col["mice"][0]]})
        _, exp = self.create(self.m, from_cage_id=col["cage"])
        self.assertEqual(rows("select mouse_id_fk from experiment_mice where experiment_id_fk=?", exp), [(col["mice"][1],)])

    def test_unknown_seed_cage_is_reported_and_the_experiment_starts_empty(self):
        cage = uniq("nocage")
        name = uniq("Exp ")
        r = self.post(self.m, "/colony/experiments/create", {"name": name, "from_cage_id": cage})
        self.assertFlash(r, f"There is no cage {cage}", "error")
        exp = one("select id from experiments where name=?", name)
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 0)


class ExperimentMembersTests(AppTestCase):

    def new_experiment(self, client):
        name = uniq("Exp ")
        client.post("/colony/experiments/create", data={"name": name})
        return one("select id from experiments where name=?", name)

    def test_add_cage_adds_its_living_mice_once(self):
        exp = self.new_experiment(self.m)
        col = self.make_colony(self.m, self.member, n_mice=2)
        r = self.post(self.m, f"/colony/experiments/{exp}/add-cage", {"cage_id": col["cage"]})
        self.assertFlash(r, "Added 2 mice", "success")
        r = self.post(self.m, f"/colony/experiments/{exp}/add-cage", {"cage_id": col["cage"]})
        self.assertFlash(r, "no living mice that are not already in the experiment", "info")
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 2)

    def test_add_cage_skips_someone_elses_mice(self):
        exp = self.new_experiment(self.m)
        col = self.make_colony(self.o, self.other, n_mice=2, purpose="Experiments")
        r = self.post(self.m, f"/colony/experiments/{exp}/add-cage", {"cage_id": col["cage"]})
        self.assertFlash(r, "2 skipped", "error")
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 0)

    def test_add_cage_reports_a_blank_or_unknown_cage(self):
        exp = self.new_experiment(self.m)
        self.assertFlash(self.post(self.m, f"/colony/experiments/{exp}/add-cage", {"cage_id": ""}),
                         "Enter a cage", "error")
        code = uniq("nocage")
        self.assertFlash(self.post(self.m, f"/colony/experiments/{exp}/add-cage", {"cage_id": code}),
                         f"Cage '{code}' not found", "error")

    def test_add_mouse_with_a_group_and_not_twice(self):
        exp = self.new_experiment(self.m)
        mouse = self.make_mouse(self.m, self.member, cage=uniq("C"))
        for _ in range(2):
            self.m.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse, "treatment_group": "vehicle"})
        self.assertEqual(rows("select mouse_id_fk, treatment_group from experiment_mice where experiment_id_fk=?", exp),
                         [(mouse, "vehicle")])

    def test_add_mouse_refuses_someone_elses_mouse(self):
        exp = self.new_experiment(self.m)
        theirs = self.make_mouse(self.o, self.other, cage=uniq("C"))
        r = self.post(self.m, f"/colony/experiments/{exp}/add-mouse", {"mouse_row_id": theirs})
        self.assertFlash(r, self.other, "error")
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 0)

    def test_member_update_sets_and_clears_the_group(self):
        exp = self.new_experiment(self.m)
        mouse = self.make_mouse(self.m, self.member, cage=uniq("C"))
        self.m.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse, "treatment_group": "drug"})
        em = one("select id from experiment_mice where experiment_id_fk=?", exp)
        self.assertSaved(self.m.post(f"/colony/experiments/{exp}/members/{em}/update", data={"note": "n1"}))
        self.assertEqual(row("select treatment_group, note from experiment_mice where id=?", em), ("drug", "n1"))
        self.assertSaved(self.m.post(f"/colony/experiments/{exp}/members/{em}/update", data={"treatment_group": ""}))
        self.assertEqual(one("select treatment_group from experiment_mice where id=?", em), "")

    def test_owner_removes_a_member_and_the_mouse_is_unchanged(self):
        exp = self.new_experiment(self.m)
        mouse = self.make_mouse(self.m, self.member, cage=uniq("C"))
        self.m.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse})
        em = one("select id from experiment_mice where experiment_id_fk=?", exp)
        self.m.post(f"/colony/experiments/{exp}/members/{em}/remove")
        self.assertEqual(count("experiment_mice", "id=?", em), 0)
        self.assertEqual(count("mice", "id=?", mouse), 1)

    def test_weights_are_only_for_mice_the_user_may_edit(self):
        mine = self.make_mouse(self.m, self.member, cage=uniq("C"))
        theirs = self.make_mouse(self.o, self.other, cage=uniq("C"))
        self.assertEqual(self.m.post(f"/colony/mice/{theirs}/weights/create", data={"grams": "20"}).status_code, 403)
        self.assertEqual(self.m.post(f"/colony/mice/{mine}/weights/create", data={"grams": "abc"}).status_code, 400)
        # Same day twice: one row, the later weight.
        for grams in ("21.5", "22.0"):
            self.m.post(f"/colony/mice/{mine}/weights/create", data={"grams": grams, "weigh_date": days_ago(1)},
                        headers=AUTOSAVE)
        self.assertEqual(rows("select weigh_date, grams from mouse_weights where mouse_id_fk=?", mine), [(days_ago(1), 22.0)])
        self.assertEqual(count("mouse_weights", "mouse_id_fk=?", theirs), 0)


class ExperimentPermissionTests(AppTestCase):
    """Another member's experiment is read only; its owner and admins edit it."""

    def setUp(self):
        name = uniq("Exp ")
        self.name = name
        self.a.post("/colony/experiments/create", data={"name": name})
        self.exp = one("select id from experiments where name=?", name)
        self.admins_mouse = self.make_mouse(self.a, self.admin, cage=uniq("C"))
        self.a.post(f"/colony/experiments/{self.exp}/add-mouse",
                    data={"mouse_row_id": self.admins_mouse, "treatment_group": "g1"})
        self.em = one("select id from experiment_mice where experiment_id_fk=?", self.exp)

    def members(self):
        return rows("select id, mouse_id_fk, treatment_group from experiment_mice where experiment_id_fk=? order by id", self.exp)

    def test_member_cannot_rename_someone_elses_experiment(self):
        r = self.autosave(self.m, f"/colony/experiments/{self.exp}/update", {"name": "renamed"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(one("select name from experiments where id=?", self.exp), self.name)

    def test_member_cannot_add_cages_or_mice_to_someone_elses_experiment(self):
        before = self.members()
        col = self.make_colony(self.m, self.member, n_mice=1)
        self.m.post(f"/colony/experiments/{self.exp}/add-cage", data={"cage_id": col["cage"]})
        self.m.post(f"/colony/experiments/{self.exp}/add-mouse", data={"mouse_row_id": col["mice"][0]})
        self.assertEqual(self.members(), before)

    def test_member_cannot_change_or_remove_members_of_someone_elses_experiment(self):
        before = self.members()
        r = self.m.post(f"/colony/experiments/{self.exp}/members/{self.em}/update", data={"treatment_group": "x"})
        self.assertEqual(r.status_code, 403)
        self.m.post(f"/colony/experiments/{self.exp}/members/{self.em}/remove")
        self.assertEqual(self.members(), before)

    def test_member_cannot_delete_someone_elses_experiment(self):
        r = self.post(self.m, f"/colony/experiments/{self.exp}/delete")
        self.assertFlash(r, f"{self.admin}’s", "error")
        self.assertEqual(count("experiments", "id=?", self.exp), 1)

    def test_detail_page_is_read_only_for_a_non_owner(self):
        html = self.get_ok(self.m, f"/colony/experiments/{self.exp}")
        self.assertRegex(html, r'name="name"[^>]*disabled')
        self.assertNotIn("Delete experiment", html)
        self.assertNotIn('data-xp-open="add"', html)

    def test_detail_page_is_editable_for_the_owner(self):
        html = self.get_ok(self.a, f"/colony/experiments/{self.exp}")
        self.assertNotRegex(html, r'name="name"[^>]*disabled')
        self.assertIn("Delete experiment", html)
        self.assertIn('data-xp-open="add"', html)

    def test_experiments_tab_shows_a_lock_on_someone_elses_experiment(self):
        html = self.get_ok(self.m, "/colony?view=experiments")
        card = row_html(html, f'/colony/experiments/{self.exp}"', "</a>")
        self.assertIn("sheet-lock", card)

    def test_admin_may_edit_a_members_experiment(self):
        name = uniq("Exp ")
        self.m.post("/colony/experiments/create", data={"name": name})
        exp = one("select id from experiments where name=?", name)
        self.assertSaved(self.autosave(self.a, f"/colony/experiments/{exp}/update", {"status": "paused"}))
        self.assertEqual(one("select status from experiments where id=?", exp), "paused")

    def test_admin_may_delete_a_members_experiment_leaving_its_mice(self):
        name = uniq("Exp ")
        self.m.post("/colony/experiments/create", data={"name": name})
        exp = one("select id from experiments where name=?", name)
        mouse = self.make_mouse(self.m, self.member, cage=uniq("C"))
        self.m.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse})
        r = self.post(self.a, f"/colony/experiments/{exp}/delete")
        self.assertFlash(r, f"Deleted experiment {name}", "success")
        self.assertEqual((count("experiments", "id=?", exp), count("experiment_mice", "experiment_id_fk=?", exp),
                          count("mice", "id=?", mouse)), (0, 0, 1))

    def test_owner_updates_fields_and_dates(self):
        r = self.autosave(self.a, f"/colony/experiments/{self.exp}/update", {
            "name": self.name + " v2", "description": "desc", "status": "done",
            "start_date": days_ago(10), "end_date": days_ago(1)})
        self.assertSaved(r)
        self.assertEqual(row("select name, description, status, start_date, end_date from experiments where id=?", self.exp),
                         (self.name + " v2", "desc", "done", days_ago(10), days_ago(1)))

    def test_end_date_before_start_date_is_refused(self):
        r = self.autosave(self.a, f"/colony/experiments/{self.exp}/update",
                          {"start_date": days_ago(1), "end_date": days_ago(10)})
        self.assertRefused(r)
        self.assertIn("before the start date", r.get_json()["error"])
        self.assertEqual(row("select start_date, end_date from experiments where id=?", self.exp), (None, None))

    def test_unknown_status_is_refused(self):
        self.assertRefused(self.autosave(self.a, f"/colony/experiments/{self.exp}/update", {"status": "exploded"}))
        self.assertEqual(one("select status from experiments where id=?", self.exp), "active")


class BulkAddToExperimentTests(AppTestCase):

    def test_adds_editable_mice_in_one_group_and_reports_the_rest(self):
        name = uniq("Exp ")
        self.m.post("/colony/experiments/create", data={"name": name})
        exp = one("select id from experiments where name=?", name)
        mine = self.make_colony(self.m, self.member, n_mice=2)["mice"]
        theirs = self.make_mouse(self.o, self.other, cage=uniq("C"))
        self.m.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mine[0]})
        r = self.post(self.m, "/colony/mice/bulk-experiment",
                      {"experiment_id": exp, "treatment_group": "vehicle", "selected_ids": [*mine, theirs]})
        text = flash_text(r)
        self.assertIn(f"Added 1 mouse to {name} as “vehicle”", text)
        self.assertIn("(1 already in it)", text)
        self.assertIn("1 skipped", text)
        self.assertEqual(rows("select mouse_id_fk, treatment_group from experiment_mice where experiment_id_fk=? order by id", exp),
                         [(mine[0], ""), (mine[1], "vehicle")])

    def test_is_one_batch_that_undo_removes(self):
        name = uniq("Exp ")
        self.m.post("/colony/experiments/create", data={"name": name})
        exp = one("select id from experiments where name=?", name)
        mice = self.make_colony(self.m, self.member, n_mice=2)["mice"]
        self.m.post("/colony/mice/bulk-experiment", data={"experiment_id": exp, "treatment_group": "", "selected_ids": mice})
        batch = row("select id, record_count, actor from batches where description=?", f"add to experiment {name}")
        self.assertEqual(batch[1:], (2, self.member))
        self.m.post(f"/batches/{batch[0]}/undo")
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 0)

# ====================================================================== strains

class StrainTests(AppTestCase):

    def make_strain(self, client, name=None, **fields):
        name = name or uniq("Strain")
        client.post("/colony/strains/create", data={"strain_name": name, **fields})
        found = one("select id from strains where strain_name=?", name)
        self.assertTrue(found, f"strain {name} was not created")
        return found

    def test_create_records_who_added_it(self):
        sid = self.make_strain(self.m, strain_number="007914", supplier="JAX")
        self.assertEqual(row("select created_by, strain_number, supplier from strains where id=?", sid),
                         (self.member, "007914", "JAX"))

    def test_duplicate_name_in_any_case_and_blank_name_are_refused(self):
        name = uniq("Strain")
        self.make_strain(self.a, name)
        self.assertFlash(self.post(self.m, "/colony/strains/create", {"strain_name": name.upper()}),
                         "already a strain", "error")
        self.assertEqual(count("strains", "lower(strain_name)=lower(?)", name), 1)
        self.assertFlash(self.post(self.m, "/colony/strains/create", {"strain_name": "  "}), "needs a name", "error")

    def test_member_renames_and_deletes_their_own_strain(self):
        sid = self.make_strain(self.m)
        new = uniq("Strain")
        self.assertSaved(self.autosave(self.m, f"/colony/strains/{sid}/update", {"strain_name": new, "description": "d"}))
        self.assertEqual(row("select strain_name, description from strains where id=?", sid), (new, "d"))
        self.m.post(f"/colony/strains/{sid}/delete")
        self.assertEqual(count("strains", "id=?", sid), 0)

    def test_member_cannot_rename_or_delete_someone_elses_strain(self):
        name = uniq("Strain")
        sid = self.make_strain(self.o, name)
        r = self.autosave(self.m, f"/colony/strains/{sid}/update", {"strain_name": "hijack"})
        self.assertRefused(r)
        self.assertIn(f"added by {self.other}", r.get_json()["error"])
        self.assertFlash(self.post(self.m, f"/colony/strains/{sid}/delete"), "only they or an admin", "error")
        self.assertEqual(one("select strain_name from strains where id=?", sid), name)

    def test_strains_from_before_creators_are_admin_only(self):
        sid = self.make_strain(self.a)
        execute("update strains set created_by='' where id=?", sid)
        r = self.autosave(self.m, f"/colony/strains/{sid}/update", {"description": "member"})
        self.assertRefused(r)
        self.assertIn("predates recorded creators", r.get_json()["error"])
        self.assertSaved(self.autosave(self.a, f"/colony/strains/{sid}/update", {"description": "admin"}))
        self.assertEqual(one("select description from strains where id=?", sid), "admin")

    def test_admin_renames_and_deletes_a_members_strain(self):
        sid = self.make_strain(self.m)
        new = uniq("Strain")
        self.assertSaved(self.autosave(self.a, f"/colony/strains/{sid}/update", {"strain_name": new}))
        self.assertEqual(one("select strain_name from strains where id=?", sid), new)
        self.assertFlash(self.post(self.a, f"/colony/strains/{sid}/delete"), f"Removed strain {new}", "success")
        self.assertEqual(count("strains", "id=?", sid), 0)

    def test_rename_onto_an_existing_name_is_refused(self):
        taken = uniq("Strain")
        self.make_strain(self.m, taken)
        sid = self.make_strain(self.m)
        r = self.autosave(self.m, f"/colony/strains/{sid}/update", {"strain_name": taken.lower()})
        self.assertRefused(r)
        self.assertIn("already a strain", r.get_json()["error"])
        self.assertRefused(self.autosave(self.m, f"/colony/strains/{sid}/update", {"strain_name": ""}))

    def test_strains_tab_locks_other_peoples_strains(self):
        theirs = self.make_strain(self.o)
        mine = self.make_strain(self.m)
        html = self.get_ok(self.m, "/colony?view=strains")
        locked = row_html(html, f'<tr data-id="{theirs}" data-created_by')
        open_ = row_html(html, f'<tr data-id="{mine}" data-created_by')
        self.assertIn("sheet-lock", locked)
        self.assertRegex(locked, r'name="strain_name"[^>]*disabled')
        self.assertNotIn("/delete", locked)
        self.assertNotIn("sheet-lock", open_)
        self.assertIn("/delete", open_)


# ====================================================================== presets

class PresetTests(AppTestCase):

    def make_option(self, value=None, field="purpose"):
        value = value or uniq("preset ")
        self.a.post("/colony/options/create", data={"field_name": field, "option_value": value})
        found = one("select id from dropdown_options where field_name=? and option_value=?", field, value)
        self.assertTrue(found, f"preset {value} was not created")
        return found

    def test_member_cannot_add_rename_or_remove_presets(self):
        value = uniq("preset ")
        r = self.post(self.m, "/colony/options/create", {"field_name": "purpose", "option_value": value})
        self.assertFlash(r, "Only an admin", "error")
        self.assertEqual(count("dropdown_options", "option_value=?", value), 0)
        oid = self.make_option()
        before = one("select option_value from dropdown_options where id=?", oid)
        self.assertRefused(self.autosave(self.m, f"/colony/options/{oid}/update", {"option_value": "renamed"}))
        self.m.post(f"/colony/options/{oid}/delete")
        self.assertEqual(one("select option_value from dropdown_options where id=?", oid), before)

    def test_admin_adds_renames_and_removes_a_preset(self):
        oid = self.make_option()
        new = uniq("preset ")
        self.assertSaved(self.autosave(self.a, f"/colony/options/{oid}/update", {"option_value": new}))
        self.assertEqual(one("select option_value from dropdown_options where id=?", oid), new)
        self.assertFlash(self.post(self.a, f"/colony/options/{oid}/delete"), "Removed the purpose preset", "success")
        self.assertEqual(count("dropdown_options", "id=?", oid), 0)

    def test_duplicate_preset_in_any_case_is_refused(self):
        value = uniq("Preset ")
        self.make_option(value)
        r = self.post(self.a, "/colony/options/create", {"field_name": "purpose", "option_value": value.lower()})
        self.assertFlash(r, "already a purpose preset", "error")
        self.assertEqual(count("dropdown_options", "field_name='purpose' and lower(option_value)=lower(?)", value), 1)

    def test_rename_onto_another_preset_or_to_blank_is_refused(self):
        taken = uniq("preset ")
        self.make_option(taken)
        oid = self.make_option()
        before = one("select option_value from dropdown_options where id=?", oid)
        self.assertRefused(self.autosave(self.a, f"/colony/options/{oid}/update", {"option_value": taken.upper()}))
        self.assertRefused(self.autosave(self.a, f"/colony/options/{oid}/update", {"option_value": " "}))
        self.assertEqual(one("select option_value from dropdown_options where id=?", oid), before)

    def test_presets_tab_is_read_only_for_members(self):
        member_html = self.get_ok(self.m, "/colony?view=settings")
        self.assertIn("Only an admin can add, rename or remove choices", member_html)
        self.assertNotIn("Add choice", member_html)
        admin_html = self.get_ok(self.a, "/colony?view=settings")
        self.assertIn("Add choice", admin_html)
        self.assertNotIn("Only an admin can add, rename or remove choices", admin_html)


# ============================================================ Add many (batch)

class BatchMiceTests(AppTestCase):

    def grid(self, specs):
        """Preview-grid form fields for rows of {field: value}."""
        return {f"rows-{i}-{key}": value for i, spec in enumerate(specs) for key, value in spec.items()}

    def test_setup_page_renders(self):
        self.assertIn("Transgene 1", self.get_ok(self.m, "/colony/mice/batch"))

    def test_preview_splits_the_prototype_by_sex_and_writes_nothing(self):
        mice = count("mice")
        r = self.m.post("/colony/mice/batch/preview", data={
            "count_female": "2", "count_male": "1", "owner": self.member, "status": "experiment"})
        self.assertEqual(r.status_code, 200)
        genders = re.findall(r'name="rows-(\d+)-gender" value="([^"]*)"', r.get_data(as_text=True))
        self.assertEqual(genders, [("0", "F"), ("1", "F"), ("2", "M")])
        self.assertEqual(count("mice"), mice)

    def test_empty_preview_or_create_is_reported(self):
        self.assertFlash(self.post(self.m, "/colony/mice/batch/preview", {"source": "grid"}), "Nothing to preview", "error")
        self.assertFlash(self.post(self.m, "/colony/mice/batch/create", {}), "Nothing to create", "error")

    def test_create_rows_marked_new_share_one_new_cage_owned_by_the_creator(self):
        mice_before = one("select coalesce(max(id), 0) from mice")
        spec = {"cage_id": "new", "owner": self.member, "status": "experiment", "transgene_1": uniq("Tg")}
        r = self.post(self.m, "/colony/mice/batch/create",
                      self.grid([{**spec, "gender": "F"}, {**spec, "gender": "F"}, {**spec, "gender": "M"}]))
        self.assertFlash(r, "Created 3 mice", "success")
        made = rows("select mouse_id, cage_id_fk, gender, transgene_1 from mice where id>? and transgene_1=? order by id",
                    mice_before, spec["transgene_1"])
        self.assertEqual([m[2] for m in made], ["F", "F", "M"])
        self.assertEqual(len({m[1] for m in made}), 1)
        self.assertEqual(one("select owner from mouse_cages where id=?", made[0][1]), self.member)
        numbers = [m[0] for m in made]
        self.assertEqual(numbers, list(range(numbers[0], numbers[0] + 3)))

    def test_create_is_one_batch_whose_undo_removes_mice_and_new_cage(self):
        tg = uniq("Tg")
        self.m.post("/colony/mice/batch/create", data=self.grid(
            [{"gender": "F", "cage_id": "new", "owner": self.member, "transgene_1": tg}] * 2))
        cage = one("select cage_id_fk from mice where transgene_1=?", tg)
        batch = row("select id, record_count from batches where actor=? and description='add 2 mice in bulk' "
                    "order by id desc limit 1", self.member)
        self.assertEqual(batch[1], 2)
        self.m.post(f"/batches/{batch[0]}/undo")
        self.assertEqual((count("mice", "transgene_1=?", tg), count("mouse_cages", "id=?", cage)), (0, 0))

    def test_dropped_rows_are_not_created(self):
        tg = uniq("Tg")
        data = self.grid([{"gender": "F", "transgene_1": tg, "owner": self.member},
                          {"gender": "M", "transgene_1": tg, "owner": self.member}])
        data["rows-1-drop"] = "1"
        self.m.post("/colony/mice/batch/create", data=data)
        self.assertEqual(rows("select gender from mice where transgene_1=?", tg), [("F",)])

    def test_a_date_of_birth_without_a_litter_is_kept(self):
        """A spreadsheet with a dob column and no litter: every mouse keeps
        its date, the ones born the same day in one automatic litter."""
        tg = uniq("Tg")
        spec = {"owner": self.member, "status": "experiment", "transgene_1": tg}
        self.m.post("/colony/mice/batch/create", data=self.grid([
            {**spec, "gender": "F", "date_of_birth": days_ago(40)},
            {**spec, "gender": "M", "date_of_birth": days_ago(40)},
            {**spec, "gender": "F", "date_of_birth": days_ago(10)}]))
        made = rows("select l.date_of_birth, l.id from mice m join litters l on l.id=m.litter_id_fk "
                    "where m.transgene_1=? order by m.id", tg)
        self.assertEqual([str(d) for d, _ in made], [days_ago(40), days_ago(40), days_ago(10)])
        self.assertEqual(made[0][1], made[1][1])       # born the same day: one litter
        self.assertNotEqual(made[0][1], made[2][1])

    def test_a_genotype_column_without_transgenes_becomes_transgene_1(self):
        note = uniq("n")
        self.m.post("/colony/mice/batch/create", data=self.grid(
            [{"gender": "F", "owner": self.member, "genotype": "C57BL/6J", "note": note}]))
        self.assertEqual(row("select transgene_1, genotype from mice where note=?", note), ("C57BL/6J", "C57BL/6J"))

    def test_a_cage_typed_in_belongs_to_whoever_made_it(self):
        cage = uniq("C")
        self.m.post("/colony/mice/batch/create", data=self.grid(
            [{"gender": "F", "cage_id": cage, "owner": self.member}]))
        self.assertEqual(one("select owner from mouse_cages where cage_id=?", cage), self.member)

    def test_new_with_both_sexes_gives_each_sex_its_own_cage(self):
        r = self.m.post("/colony/mice/batch/preview", data={
            "count_female": "2", "count_male": "1", "cage_id": "new", "owner": self.member})
        cages = re.findall(r'name="rows-\d+-cage_id" value="([^"]*)"', r.get_data(as_text=True))
        self.assertEqual(cages, ["new-F", "new-F", "new-M"])
        tg = uniq("Tg")
        self.m.post("/colony/mice/batch/create", data=self.grid(
            [{"gender": g, "cage_id": c, "owner": self.member, "transgene_1": tg}
             for g, c in (("F", "new-F"), ("F", "new-F"), ("M", "new-M"))]))
        made = rows("select gender, cage_id_fk from mice where transgene_1=? order by id", tg)
        self.assertEqual(made[0][1], made[1][1])
        self.assertNotEqual(made[0][1], made[2][1])

    def test_a_breeding_cage_may_be_asked_for(self):
        r = self.m.post("/colony/mice/batch/preview", data={
            "count_female": "1", "count_male": "1", "cage_id": "new", "one_new_cage": "1", "owner": self.member})
        html = r.get_data(as_text=True)
        self.assertEqual(re.findall(r'name="rows-\d+-cage_id" value="([^"]*)"', html), ["new", "new"])
        self.assertIn("would get females and males together", html)

    def test_the_template_has_the_headers_the_page_lists(self):
        r = self.m.get("/colony/mice/batch/template.csv")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["Content-Disposition"])
        header = r.get_data(as_text=True).lstrip("\ufeff").splitlines()[0]
        self.assertTrue(header.startswith("sex,transgene_1"))
        self.assertIn(header.replace(",", " · "), self.get_ok(self.m, "/colony/mice/batch"))

    def csv_preview(self, text):
        return self.m.post("/colony/mice/batch/preview", content_type="multipart/form-data", data={
            "file": (io.BytesIO(text.encode()), "mice.csv")}).get_data(as_text=True)

    def dobs(self, html):
        return re.findall(r'name="rows-\d+-date_of_birth" value="([^"]*)"', html)

    def test_excel_dates_in_a_csv_are_read(self):
        html = self.csv_preview("sex,dob\nF,3/14/2026\nM,2026-02-01\nF,4/5/26\n")
        self.assertEqual(self.dobs(html), ["2026-03-14", "2026-02-01", "2026-04-05"])
        self.assertNotIn("read month first", html)     # 3/14 settles the order

    def test_day_first_dates_are_recognised_from_the_file(self):
        html = self.csv_preview("sex,dob\nF,14/03/2026\nM,04.05.2026\n")
        self.assertEqual(self.dobs(html), ["2026-03-14", "2026-05-04"])

    def test_ambiguous_and_unreadable_dates_are_reported(self):
        html = self.csv_preview("sex,dob\nF,03/04/2026\nM,soon\n")
        self.assertEqual(self.dobs(html), ["2026-03-04", ""])
        self.assertIn("read month first", html)
        self.assertIn("Row 3: “soon” is not a date", html)


# ================================================================ global search

class GlobalSearchTests(AppTestCase):

    def results(self, client, q):
        r = client.get("/search", query_string={"q": q})
        self.assertEqual(r.status_code, 200)
        return r.get_json()["results"]

    def test_empty_query_returns_nothing(self):
        self.assertEqual(self.results(self.m, "  "), [])

    def test_finds_a_mouse_by_number_linking_to_everyones_sheet(self):
        mouse = self.make_mouse(self.o, self.other, cage=uniq("C"))
        number = mouse_number(mouse)
        hits = [x for x in self.results(self.m, str(number)) if x["type"] == "mouse"]
        self.assertEqual([x["id"] for x in hits], [number])
        self.assertEqual(hits[0]["url"], f"/colony?view=mice&scope=all&q={number}")

    def test_finds_cages_litters_experiments_and_strains(self):
        tag = uniq("zq")
        self.make_cage(self.o, f"{tag}cage")
        self.make_litter(self.o, f"{tag}lit")
        self.o.post("/colony/experiments/create", data={"name": f"{tag} exp"})
        self.o.post("/colony/strains/create", data={"strain_name": f"{tag}strain"})
        exp = one("select id from experiments where name=?", f"{tag} exp")
        found = {(x["type"], x["label"], x["url"]) for x in self.results(self.m, tag)}
        self.assertEqual(found, {
            ("cage", f"Cage {tag}cage", f"/colony?view=cages&scope=all&q={tag}cage"),
            ("litter", f"Litter {tag}lit", f"/colony?view=litters&q={tag}lit"),
            ("experiment", f"{tag} exp", f"/colony/experiments/{exp}"),
            ("strain", f"{tag}strain", f"/colony?view=strains&q={tag}strain"),
        })

    def test_notebook_pages_are_only_found_by_their_owner(self):
        tag = uniq("zqpage")
        tab = self.o.post("/notebook/tabs/create", data={"title": "T"}).get_json()["id"]
        self.o.post("/notebook/pages/create", data={"tab_id": tab, "title": f"{tag} notes"})
        self.assertEqual([x["label"] for x in self.results(self.o, tag) if x["type"] == "page"], [f"{tag} notes"])
        self.assertEqual([x for x in self.results(self.m, tag) if x["type"] == "page"], [])


# ============================================================ admin overview

class AdminColonyOverviewTests(AppTestCase):

    def test_the_overview_is_settings_statistics_now(self):
        self.make_colony(self.m, self.member, n_mice=2)
        r = self.a.get("/admin/colony")
        self.assertTrue(r.headers["Location"].endswith("/settings#stats"))
        html = self.get_ok(self.a, "/settings")
        stats = html.split('id="stats"', 1)[1].split('id="general"', 1)[0]
        self.assertIn("living mice", stats)
        self.assertIn(self.member, stats)

    def test_members_are_turned_away(self):
        r = self.m.get("/admin/colony", follow_redirects=True)
        self.assertFlash(r, "Admin access required", "error")


if __name__ == "__main__":
    unittest.main()
