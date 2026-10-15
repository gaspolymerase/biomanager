"""Experiments on any database's animals: mice, zebrafish, flies, worms and
the lab's own organisms.

An experiment is on one database's animals (Experiment.db): the mouse
colony's mice (ExperimentMouse, as before), or zebrafish rows, fly vials,
worm plates or organisms (ExperimentSubject). Each has

- a readout, measured over its days: a mouse's body weight (MouseWeight,
  so it is on the mouse as well), or a length, a score, how many are alive
  of those at the start (ExperimentReading). The readout fits the animal:
  flies and worms have no body weight, so their usual one is survival, and
  a zebrafish's is survival too (length and a phenotype are there);
- manipulations, planned and recorded (app/experiment_steps.py), with the
  kinds the animal has: injections for a mouse, drug in the water or a heat
  shock for fish, drug in the food or a temperature shift for flies, RNAi
  feeding for worms.

The page (templates/experiment.html, static/experiment-page.js) is the
same for all: the experiment's details; a sheet of its animals that
switches between the treatments each got and the readout, editable in
place; the button that records a manipulation; and the readout over the
days as a chart, the manipulation days dashed, each saying what was done
when you point at it.
"""
from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from flask import Blueprint, abort, flash, g, jsonify, redirect, render_template, request, url_for
from sqlalchemy import func, select

from . import access, derived, i18n, lab, positions
from .db import SessionLocal
from .i18n import gettext, ngettext
from .models import (CageRecord, ClutchRecord, Experiment, ExperimentMouse, ExperimentReading, ExperimentSubject,
                     FishRecord, MouseRecord, MouseWeight, OrgCohort, Organism, StockUnit, TankRecord)

bp = Blueprint("experiments", __name__, url_prefix="/experiments")

# "planned": designed, not started; nothing of it is due until Start (app/experiment_plan.py).
STATUSES = ["planned", "active", "paused", "done", "cancelled"]
STARTED = STATUSES[1:]

# ---------------------------------------------------------------- readouts

READOUTS = {
    "body_weight": {"label": "Body weight", "unit": "g", "kind": "value", "verb": "Weigh"},
    "tumour_volume": {"label": "Tumour volume", "unit": "mm³", "kind": "value", "verb": "Measure"},
    "score": {"label": "Clinical score", "unit": "", "kind": "value", "verb": "Score"},
    "length": {"label": "Standard length", "unit": "mm", "kind": "value", "verb": "Measure"},
    "survival": {"label": "Survival", "unit": "alive", "kind": "fraction", "verb": "Count"},
    "phenotype": {"label": "Phenotype", "unit": "affected", "kind": "fraction", "verb": "Score"},
    "eclosion": {"label": "Eclosed adults", "unit": "flies", "kind": "value", "verb": "Count"},
    "climbing": {"label": "Climbing", "unit": "climbed", "kind": "fraction", "verb": "Test"},
    "brood": {"label": "Brood size", "unit": "progeny", "kind": "value", "verb": "Count"},
    "paralysis": {"label": "Paralysis", "unit": "paralysed", "kind": "fraction", "verb": "Score"},
}
READOUT_CHOICES = {
    "mouse": ["body_weight", "tumour_volume", "score", "survival"],
    "fish": ["survival", "length", "phenotype"],
    "fly": ["survival", "eclosion", "climbing"],
    "worm": ["survival", "brood", "paralysis"],
    "organism": ["body_weight", "length", "survival", "score"],
}

# What can be done to each kind of animal: key → (label, icon). "reading"
# is a day of the readout ("weigh" in the first version of this).
KINDS = {
    "mouse": [("injection", "Injection", "syringe"), ("challenge", "Challenge", "virus"),
              ("treatment", "Treatment", "droplet"), ("surgery", "Surgery", "stethoscope"),
              ("behavior", "Behaviour", "gauge"), ("stimulation", "Stimulation", "bolt"), ("sample", "Sample", "vial"),
              ("imaging", "Imaging", "microscope")],
    "fish": [("immersion", "Drug in the water", "droplet"), ("microinjection", "Microinjection", "syringe"),
             ("heat_shock", "Heat shock", "temperature"), ("injury", "Injury or fin clip", "stethoscope"),
             ("feeding", "Feeding", "flask"), ("behavior", "Behaviour", "gauge"), ("stimulation", "Stimulation", "bolt"),
             ("imaging", "Imaging", "microscope"), ("sample", "Sample", "vial")],
    "fly": [("food", "Drug in the food", "flask"), ("temperature", "Temperature shift", "temperature"),
            ("starvation", "Starvation", "clock"), ("infection", "Infection", "virus"),
            ("irradiation", "Irradiation", "bolt"), ("transfer", "Flip to fresh food", "repeat"),
            ("behavior", "Behaviour", "gauge")],
    "worm": [("plate", "Drug on the plate", "flask"), ("rnai", "RNAi feeding", "dna"),
             ("temperature", "Temperature shift", "temperature"), ("heat_shock", "Heat shock", "temperature"),
             ("starvation", "Starvation", "clock"), ("infection", "Infection", "virus"),
             ("transfer", "Transfer to a fresh plate", "repeat")],
    "organism": [("injection", "Injection", "syringe"), ("treatment", "Treatment", "droplet"),
                 ("challenge", "Challenge", "virus"), ("surgery", "Surgery", "stethoscope"),
                 ("behavior", "Behaviour", "gauge"), ("stimulation", "Stimulation", "bolt"), ("sample", "Sample", "vial"),
                 ("imaging", "Imaging", "microscope")],
}
READING_KIND = ("reading", "Readout day", "chart")
OTHER_KIND = ("other", "Other", "note")


def kinds_for(family: str) -> list[tuple[str, str, str]]:
    return [*KINDS.get(family, KINDS["organism"]), READING_KIND, OTHER_KIND]


def kind_info(family: str, key: str) -> tuple[str, str]:
    """(label, icon) of a manipulation kind; any kind a step already has
    keeps a label, whichever animal it is on."""
    if key == "weigh":
        key = "reading"
    for k, label, icon in kinds_for(family):
        if k == key:
            return label, icon
    for rows in KINDS.values():
        for k, label, icon in rows:
            if k == key:
                return label, icon
    return key.replace("_", " ").capitalize() or "Other", "note"


class NotYours(str):
    """set_reading's answer when the animal's own record is not this
    person's to change: skipped, unlike a value that is wrong."""


def is_reading(kind: str) -> bool:
    return kind in ("reading", "weigh")


# ---------------------------------------------------------------- databases

@dataclass
class Place:
    """The database an experiment is on, as far as experiments care."""
    key: str                 # colony, zebrafish, stocks:<key>, organisms:<key>
    label: str               # Mouse colony, Zebrafish, Drosophila…
    family: str              # mouse | fish | fly | worm | organism
    subject_kind: str        # mouse | fish | unit | organism
    noun: str
    nouns: str
    housing: str             # cage, tank, rack, housing noun
    icon: str
    list_url: str
    home_url: str = ""
    module: object = None
    mv: object = None


def place_for(session, key: str) -> Place | None:
    """None when there is no such database, or it's switched off or not
    this person's to see."""
    from . import organism_service, stock_service
    kind, _, sub = (key or "colony").partition(":")
    features = lab.features_on(session)
    if kind == "colony":
        if not features.get("colony", True):
            return None
        return Place("colony", lab.FEATURES["colony"].label, "mouse", "mouse", "mouse", "mice", "cage", "mouse",
                     url_for("colony", view="experiments"), url_for("colony"))
    if kind == "zebrafish":
        if not features.get("zebrafish", True):
            return None
        return Place("zebrafish", lab.FEATURES["zebrafish"].label, "fish", "fish", "fish row", "fish rows", "tank",
                     "fish", url_for("zebrafish", view="experiments"), url_for("zebrafish", view="fish"))
    if kind == "stocks":
        module = stock_service.get_module(session, sub)
        if module is None or not lab.can_see(module):
            return None
        mv = stock_service.view(module)
        family = "worm" if module.kind == "worm" else "fly"
        return Place(key, mv.label, family, "unit", mv.unit, mv.units, mv.rack_noun, module.icon or family,
                     url_for("stocks.module", key=module.key, view="experiments"),
                     url_for("stocks.module", key=module.key), module, mv)
    if kind == "organisms":
        module = organism_service.get_module(session, sub)
        if module is None or not lab.can_see(module):
            return None
        mv = organism_service.view(module)
        return Place(key, mv.label, "organism", "organism", mv.organism_noun, mv.organism_noun_plural,
                     mv.housing_noun, module.icon or "paw", url_for("organisms.module", key=module.key, view="experiments"),
                     url_for("organisms.module", key=module.key), module, mv)
    return None


def page_url(exp: Experiment) -> str:
    if (exp.db or "colony") == "colony":
        return url_for("experiment_detail", experiment_id=exp.id)
    return url_for("experiments.page", experiment_id=exp.id)


def readout_of(exp: Experiment, place: Place) -> dict:
    """The readout: the one chosen, or the database's usual one."""
    try:
        chosen = json.loads(exp.readout or "{}")
    except ValueError:
        chosen = {}
    if not isinstance(chosen, dict) or not chosen.get("key"):
        chosen = {"key": READOUT_CHOICES.get(place.family, ["survival"])[0]}
    preset = READOUTS.get(chosen["key"])
    if preset:
        return {"key": chosen["key"], **preset}
    if chosen["key"] == "derived" and chosen.get("formula") in derived.FORMULAS:
        return derived_readout(derived.FORMULAS[chosen["formula"]], str(chosen.get("control") or ""))
    kind = chosen.get("kind") if chosen.get("kind") in ("value", "fraction") else "value"
    return {"key": "custom", "label": str(chosen.get("label") or "Measurement")[:60],
            "unit": str(chosen.get("unit") or "")[:20], "kind": kind, "verb": "Record"}


def readout_choices(place: Place) -> list[dict]:
    return [{"key": k, **READOUTS[k]} for k in READOUT_CHOICES.get(place.family, READOUT_CHOICES["organism"])]


# Readouts BioManager works out from what is measured (app/derived.py): the
# experiment keeps {"key": "derived", "formula", "control"} in its readout,
# and each reading keeps its inputs, as JSON, in its note.
DERIVED = "derived:"


def derived_readout(f: derived.Formula, control: str = "") -> dict:
    return {"key": f"{DERIVED}{f.key}", "formula": f.key, "label": f.label, "unit": f.unit, "kind": "value",
            "verb": f.verb, "expression": f.expression, "description": f.description,
            "inputs": [i.as_dict() for i in f.inputs], "per_group": f.per_group,
            "control": control if f.per_group else ""}


def derived_choices(place: Place) -> list[dict]:
    """The formulas that fit this database's animals, as readouts to choose."""
    return [derived_readout(f) for f in derived.FORMULAS.values() if not f.families or place.family in f.families]


def new_readout_choices(family: str) -> dict:
    """The New experiment form's readouts for a family of animals
    (templates/_readout_choice.html): measured ones, and worked out."""
    from types import SimpleNamespace
    place = SimpleNamespace(family=family)
    return {"readouts": readout_choices(place), "derived": derived_choices(place)}


def readout_json(key: str, control: str = "") -> str | None:
    """What an experiment keeps for a readout chosen by its key (a built-in
    one or "derived:<formula>"); None for a key that is neither."""
    if key in READOUTS:
        return json.dumps({"key": key})
    f = derived.FORMULAS.get(key[len(DERIVED):]) if key.startswith(DERIVED) else None
    if f is None:
        return None
    chosen = {"key": "derived", "formula": f.key}
    if f.per_group and control.strip():
        chosen["control"] = control.strip()[:80]
    return json.dumps(chosen)


# Other measures. An experiment has one main readout: the one its plan,
# tests by day, chart and bench mode are about (app/experiment_plan.py
# keeps it as planned). A plate can fill another measure of the same
# animals beside it (an IL-6 ELISA in a tumour experiment): kept in the
# readout's JSON, {"measures": [{"id", "label", "unit"}]}, with readings
# keyed "measure:<id>", shown on the page's sheet and chart (a chip each),
# in Export and in the assistant's experiment_summary. A measure may be one
# BioManager works out from two Cts ({"formula": "delta_ct"}, or ΔΔCt and
# fold change with their "control" group), as the main readout can be.
MEASURE = "measure:"
MEASURE_FORMULAS = ("delta_ct", "delta_delta_ct", "fold_change")


def _chosen(exp: Experiment) -> dict:
    try:
        chosen = json.loads(exp.readout or "{}")
    except ValueError:
        chosen = {}
    return chosen if isinstance(chosen, dict) else {}


def measure_readout(ident: int, label: str, unit: str = "", formula: str = "", control: str = "") -> dict:
    out = {"key": f"{MEASURE}{ident}", "id": ident, "label": " ".join(str(label).split())[:60],
           "unit": str(unit or "").strip()[:20], "kind": "value", "verb": "Record", "measure": True}
    f = derived.FORMULAS.get(str(formula or "")) if formula in MEASURE_FORMULAS else None
    if f is not None:
        out = {**derived_readout(f, str(control or "").strip()[:80]), **out, "unit": f.unit, "verb": f.verb,
               "kind_label": f.label}
    return out


def measures_of(exp: Experiment) -> list[dict]:
    """The experiment's other measures, in the order they were added."""
    out = []
    for m in _chosen(exp).get("measures") or []:
        if isinstance(m, dict) and str(m.get("id") or "").isdigit() and str(m.get("label") or "").strip():
            out.append(measure_readout(int(m["id"]), m["label"], m.get("unit") or "", m.get("formula") or "",
                                       m.get("control") or ""))
    return out


def readout_for(exp: Experiment, place: Place, key) -> dict | None:
    """The main readout (no key, "main" or its own key), or another measure
    by its key or name; None for neither."""
    main = readout_of(exp, place)
    text = str(key or "").strip()
    if not text or text in ("main", main["key"]):
        return main
    for m in measures_of(exp):
        if text == m["key"] or text.lower() == m["label"].lower():
            return m
    return None


def _measure_name(exp: Experiment, place: Place, label, keep: str = "") -> str:
    """A measure's name, tidied; a ValueError, in words, without one or for
    the main readout's or another measure's (`keep`: the measure renamed)."""
    label = " ".join(str(label or "").split())[:60]
    if not label:
        raise ValueError(gettext("Name the measure, e.g. IL-6 ELISA."))
    main = readout_of(exp, place)
    if label.lower() in {main["label"].lower(), i18n.translate_value(main["label"]).lower()}:
        raise ValueError(gettext("%(label)s is already the main readout: choose it in the list.", label=label))
    other = readout_for(exp, place, label)
    if other is not None and keep and other["key"] != keep:
        raise ValueError(gettext("This experiment already has a measure called “%(label)s”.", label=label))
    return label


def _measure_kind(formula) -> str:
    """A measure's kind as kept: "" for a number, else one of MEASURE_FORMULAS;
    a ValueError, in words, for anything else."""
    formula = str(formula or "").strip()
    if formula in ("", "value"):
        return ""
    if formula not in MEASURE_FORMULAS:
        raise ValueError(gettext("“%(kind)s” isn't a kind of measure: a number, or one of %(formulas)s.",
                                 kind=formula, formulas=", ".join(MEASURE_FORMULAS)))
    return formula


def add_measure(exp: Experiment, place: Place, label, unit="", formula="", control="") -> dict:
    """Another measure of the experiment's animals, kept with it (the one of
    that name when it has one already); a ValueError, in words, without a
    name or for the main readout's own. `formula` makes it one BioManager
    works out from two Cts (MEASURE_FORMULAS), against `control` for ΔΔCt."""
    label = _measure_name(exp, place, label)
    formula = _measure_kind(formula)
    found = readout_for(exp, place, label)
    if found is not None:
        if (found.get("formula") or "") != formula:
            raise ValueError(gettext("This experiment already has a measure called “%(label)s”, of another kind: choose it in the list, or give the new one another name.",  # noqa: E501
                                     label=found["label"]))
        return found
    chosen = _chosen(exp)
    kept = [m for m in chosen.get("measures") or [] if isinstance(m, dict)]
    # A deleted measure's number is never given again: its plates keep it.
    last = chosen.get("last_measure") if isinstance(chosen.get("last_measure"), int) else 0
    ident = max([int(m["id"]) for m in kept if str(m.get("id") or "").isdigit()] + [last]) + 1
    entry = {"id": ident, "label": label, "unit": str(unit or "").strip()[:20]}
    if formula:
        entry["formula"] = formula
        if derived.FORMULAS[formula].per_group and str(control or "").strip():
            entry["control"] = str(control).strip()[:80]
    kept.append(entry)
    chosen["measures"], chosen["last_measure"] = kept, ident
    exp.readout = json.dumps(chosen, ensure_ascii=False)
    return measure_readout(ident, label, entry["unit"], formula, entry.get("control", ""))


def change_measure(exp: Experiment, place: Place, key: str, data: dict) -> dict:
    """A measure renamed (`label`), its unit, its kind or its control group
    changed (`unit`, `kind`, `control`); its id, and so its readings, stay
    (`rekind_readings` turns them into the new kind). A ValueError, in
    words, for a name or kind it can't have; a LookupError for no such
    measure."""
    m = readout_for(exp, place, key)
    if m is None or not m.get("measure"):
        raise LookupError(key)
    kind = _measure_kind(data["kind"]) if "kind" in data else (m.get("formula") or "")
    per_group = bool(kind) and derived.FORMULAS[kind].per_group
    chosen = _chosen(exp)
    for entry in chosen.get("measures") or []:
        if isinstance(entry, dict) and str(entry.get("id")) == str(m["id"]):
            if "label" in data:
                entry["label"] = _measure_name(exp, place, data.get("label"), keep=m["key"])
            if kind:
                entry["formula"], entry["unit"] = kind, ""
            else:
                entry.pop("formula", None)
                if "unit" in data:
                    entry["unit"] = str(data.get("unit") or "").strip()[:20]
            if "control" in data and per_group:
                entry["control"] = str(data.get("control") or "").strip()[:80]
            if not per_group:
                entry.pop("control", None)
            elif kind != (m.get("formula") or "") and not str(entry.get("control") or "").strip():
                raise ValueError(gettext("Say which group is the control of %(measure)s (its Control group): it is worked out against that group's mean ΔCt.",  # noqa: E501
                                         measure=entry["label"]))
    exp.readout = json.dumps(chosen, ensure_ascii=False)
    return readout_for(exp, place, m["key"])


def rekind_readings(session, exp: Experiment, place: Place, before: dict, after: dict) -> None:
    """A measure's readings as its new kind: a number keeps each value; ΔCt
    works each out from its two Cts again (a number becomes a ΔCt as it is);
    ΔΔCt and fold change, every day against the control group. A ValueError,
    in words, when readings have nothing the new kind can be worked out
    from (a number or a typed ΔCt has no Cts for ΔΔCt)."""
    old, new = before.get("formula") or "", after.get("formula") or ""
    if old == new:
        return
    rows = list(session.scalars(select(ExperimentReading).where(
        ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.readout_key == after["key"])))
    if not new:
        for r in rows:
            r.note = ""
        return
    f = derived.FORMULAS[new]
    bare = [r for r in rows if derived.read_note(r.note) is None]
    if bare and f.per_group:
        raise ValueError(ngettext("%(num)s reading of %(measure)s has no two Cts to work %(kind)s out from: delete it, or keep the kind.",  # noqa: E501
                                  "%(num)s readings of %(measure)s have no two Cts to work %(kind)s out from: delete them, or keep the kind.",  # noqa: E501
                                  len(bare), measure=after["label"], kind=i18n.translate_value(f.label)))
    for r in bare:
        if r.value is None or abs(r.value) > derived.CT_MAX:
            raise ValueError(gettext("A ΔCt of %(value)s isn't one a run gives (at most %(max)s cycles either way).",
                                     value=f"{r.value:g}" if r.value is not None else "", max=f"{derived.CT_MAX:g}"))
    if not f.per_group:
        for r in rows:
            inputs = derived.read_note(r.note)
            if inputs is not None:
                r.value = derived.compute(new, inputs)
        return
    subs = subjects(session, exp, place)
    session.flush()
    for on in sorted({r.read_on for r in rows}):
        _store_day(session, exp, after, on, subs)


def drop_measure(exp: Experiment, place: Place, key: str) -> dict:
    """A measure taken out of the experiment's JSON (the caller deletes its
    readings); a LookupError for no such measure."""
    m = readout_for(exp, place, key)
    if m is None or not m.get("measure"):
        raise LookupError(key)
    chosen = _chosen(exp)
    chosen["measures"] = [e for e in chosen.get("measures") or []
                          if not (isinstance(e, dict) and str(e.get("id")) == str(m["id"]))]
    last = chosen.get("last_measure") if isinstance(chosen.get("last_measure"), int) else 0
    chosen["last_measure"] = max(last, m["id"])
    exp.readout = json.dumps(chosen, ensure_ascii=False)
    return m


def keep_measures(new_json: str | None, before: str | None) -> str | None:
    """A new main readout's JSON with the other measures the experiment had."""
    try:
        old = json.loads(before or "{}")
    except ValueError:
        old = {}
    if not isinstance(old, dict) or not (old.get("measures") or old.get("last_measure")):
        return new_json
    chosen = json.loads(new_json or "{}")
    for name in ("measures", "last_measure"):
        if old.get(name):
            chosen[name] = old[name]
    return json.dumps(chosen, ensure_ascii=False)


def weighs(exp: Experiment, place: Place) -> bool:
    """Whether the readout is the animals' weight, which doses per body
    weight are worked out from."""
    return readout_of(exp, place)["key"] == "body_weight"


# ---------------------------------------------------------------- the animals

@dataclass
class Subject:
    key: str                  # "mouse:12", "fish:3", "unit:7", "organism:40"
    kind: str
    id: int
    label: str
    sex: str = ""
    genotype: str = ""
    housing: str = ""
    group: str = ""
    start: int | None = None  # animals at the start (a group); 1 for one animal
    note: str = ""
    alive: bool = True
    can_edit: bool = True     # its own record may be changed (a mouse's weight)
    row: object = None        # the membership row
    record: object = None     # the animal's own record

    @property
    def left_out(self) -> str:
        """Why it is left out of the analysis; blank when it is in."""
        return (getattr(self.row, "left_out", "") or "") if self.row is not None else ""

    @property
    def card(self) -> str:
        """The anchor its cage or housing card's QR code opens (#cage-12…),
        for scanning in bench mode."""
        rec = self.record
        if self.kind == "mouse" and getattr(rec, "cage_id_fk", None):
            return f"cage-{rec.cage_id_fk}"
        if self.kind == "fish" and getattr(rec, "tank_id_fk", None):
            return f"tank-{rec.tank_id_fk}"
        if self.kind == "organism" and getattr(rec, "housing_id_fk", None):
            return f"unit-{rec.housing_id_fk}"
        return ""

    def as_dict(self) -> dict:
        return {"key": self.key, "label": self.label, "sex": self.sex, "genotype": self.genotype,
                "housing": self.housing, "group": self.group, "start": self.start, "note": self.note,
                "alive": self.alive, "card": self.card, "left_out": self.left_out}


def subjects(session, exp: Experiment, place: Place, group: str = "") -> list[Subject]:
    """The experiment's animals, in order; only one treatment group's when
    `group` is given."""
    from .app import can_edit_mouse, mouse_is_active
    out: list[Subject] = []
    if place.subject_kind == "mouse":
        for em in exp.memberships:
            m = em.mouse
            if m is None:
                continue
            out.append(Subject(f"mouse:{m.id}", "mouse", m.id, f"#{m.mouse_id}", m.gender or "", m.genotype or "",
                               m.cage.cage_id if m.cage else "", em.treatment_group or "", 1, em.note or "",
                               mouse_is_active(m), can_edit_mouse(m), em, m))
        out.sort(key=lambda s: s.record.mouse_id or 0)
    else:
        rows = session.scalars(select(ExperimentSubject).where(ExperimentSubject.experiment_id_fk == exp.id)
                               .order_by(ExperimentSubject.id)).all()
        by_kind: dict[str, list[int]] = {}
        for r in rows:
            by_kind.setdefault(r.subject_kind, []).append(r.subject_id)
        records = {}
        if by_kind.get("fish"):
            records.update({("fish", f.id): f for f in session.scalars(select(FishRecord).where(FishRecord.id.in_(by_kind["fish"])))})
        if by_kind.get("unit"):
            records.update({("unit", u.id): u for u in session.scalars(select(StockUnit).where(StockUnit.id.in_(by_kind["unit"])))})
        if by_kind.get("organism"):
            records.update({("organism", o.id): o for o in session.scalars(select(Organism).where(Organism.id.in_(by_kind["organism"])))})
        if by_kind.get("clutch"):
            records.update({("clutch", c.id): c for c in session.scalars(select(ClutchRecord).where(ClutchRecord.id.in_(by_kind["clutch"])))})
        if by_kind.get("cohort"):
            records.update({("cohort", c.id): c for c in session.scalars(select(OrgCohort).where(OrgCohort.id.in_(by_kind["cohort"])))})
        for r in rows:
            rec = records.get((r.subject_kind, r.subject_id))
            if rec is None:
                continue
            out.append(_describe(place, r, rec))
    wanted = (group or "").strip().lower()
    if wanted:
        out = [s for s in out if s.group.strip().lower() == wanted]
    return out


def _describe(place: Place, row: ExperimentSubject, rec) -> Subject:
    key = f"{row.subject_kind}:{rec.id}"
    common = {"group": row.treatment_group or "", "note": row.note or "", "row": row, "record": rec}
    if row.subject_kind == "fish":
        tank = rec.tank.tank_id if rec.tank else ""
        label = " · ".join(filter(None, [tank, rec.individual_id or (rec.line.name if rec.line else "")]))
        return Subject(key, "fish", rec.id, label or f"Fish {rec.id}", rec.sex or "", rec.genotype or "", tank,
                       start=row.start_count if row.start_count is not None else rec.count,
                       alive=rec.status == "alive", **common)
    if row.subject_kind == "clutch":
        # Embryos or larvae of one clutch, before they are fish rows in a tank.
        start = row.start_count if row.start_count is not None else (rec.larvae_count or rec.embryo_count or None)
        return Subject(key, "clutch", rec.id, f"Clutch {rec.clutch_id}", "", rec.line.name if rec.line else "",
                       f"fertilised {rec.date_of_fertilization.isoformat()}", start=start, alive=True, **common)
    if row.subject_kind == "cohort":
        start = row.start_count if row.start_count is not None else (rec.count_initial or rec.count_current or None)
        return Subject(key, "cohort", rec.id, rec.code, "", rec.line.name if rec.line else "", rec.stage or "",
                       start=start, alive=bool(rec.count_current), **common)
    if row.subject_kind == "unit":
        mv = place.mv
        from .stock_service import position_label
        where = " ".join(filter(None, [rec.rack.name if rec.rack else "", position_label(rec)]))
        return Subject(key, "unit", rec.id, mv.code(rec) if mv else f"#{rec.number}", "", rec.genotype or "", where,
                       start=row.start_count, alive=bool(rec.active), **common)
    mv = place.mv
    label = rec.code or f"{rec.count} {mv.organism_noun_plural if mv else ''}".strip()
    alive = mv.is_alive(rec) if mv else rec.death_on is None
    return Subject(key, "organism", rec.id, label, rec.sex or "", rec.genotype or "",
                   rec.housing.code if getattr(rec, "housing", None) else "",
                   start=row.start_count if row.start_count is not None else (rec.count or 1), alive=alive, **common)


def find_subject(session, exp, place, key: str) -> Subject | None:
    return next((s for s in subjects(session, exp, place) if s.key == key), None)


def candidates(session, exp: Experiment, place: Place) -> dict:
    """What can be added: {"groups": [(value, label, count)], "single": [(value, label)]}
    — a cage, tank, rack or housing of animals, or one."""
    from .app import can_edit_mouse, mouse_is_active
    have = {s.key for s in subjects(session, exp, place)}
    groups, single = [], []
    if place.subject_kind == "mouse":
        for c in positions.natural_sorted(session.scalars(select(CageRecord)), key=lambda c: c.cage_id):
            n = sum(1 for m in c.mice if mouse_is_active(m) and f"mouse:{m.id}" not in have)
            if n:
                groups.append((c.cage_id, gettext("Cage %(cage)s", cage=c.cage_id), n))
        for m in session.scalars(select(MouseRecord).where(MouseRecord.date_of_death.is_(None)).order_by(MouseRecord.mouse_id)):
            if f"mouse:{m.id}" not in have and mouse_is_active(m) and can_edit_mouse(m):
                single.append((str(m.id), f"#{m.mouse_id} · {m.gender or '?'} · {m.genotype or gettext('no genotype')}"))
    elif place.subject_kind == "fish":
        for t in positions.natural_sorted(session.scalars(select(TankRecord).where(TankRecord.active.is_(True))),
                                          key=lambda t: t.tank_id):
            live = [f for f in t.fish if f.status == "alive" and f"fish:{f.id}" not in have]
            if live:
                groups.append((str(t.id), gettext("Tank %(tank)s", tank=t.tank_id), len(live)))
                single += [(f"fish:{f.id}", f"{t.tank_id} · {f.count} {f.sex or ''} {f.line.name if f.line else ''} {f.genotype or ''}".strip())
                           for f in live]
        recent = date.today().toordinal() - 60
        for c in session.scalars(select(ClutchRecord).order_by(ClutchRecord.date_of_fertilization.desc()).limit(200)):
            if f"clutch:{c.id}" not in have and c.date_of_fertilization.toordinal() >= recent:
                n = c.larvae_count or c.embryo_count
                single.append((f"clutch:{c.id}", " · ".join(filter(None, [
                    gettext("Clutch %(clutch)s", clutch=c.clutch_id), c.date_of_fertilization.isoformat(),
                    gettext("%(count)s larvae", count=n) if n else "", c.line.name if c.line else ""]))))
    elif place.subject_kind == "unit":
        units = session.scalars(select(StockUnit).where(StockUnit.module_id_fk == place.module.id,
                                                        StockUnit.active.is_(True)).order_by(StockUnit.number)).all()
        racks: dict = {}
        for u in units:
            if f"unit:{u.id}" in have:
                continue
            if u.rack is not None:
                racks.setdefault((u.rack.id, u.rack.name), []).append(u)
            single.append((str(u.id), f"{place.mv.code(u)} · {u.genotype or gettext('no genotype')}"))
        groups = [(str(rid), f"{i18n.translate_value(place.housing).capitalize()} {name}", len(us)) for (rid, name), us in sorted(racks.items(), key=lambda kv: positions.place_order(kv[0][1]))]
    else:
        from .models import OrgHousing
        rows = session.scalars(select(Organism).where(Organism.module_id_fk == place.module.id).order_by(Organism.id)).all()
        houses: dict = {}
        for o in rows:
            if f"organism:{o.id}" in have or not place.mv.is_alive(o):
                continue
            if o.housing_id_fk:
                houses.setdefault(o.housing_id_fk, []).append(o)
            single.append((f"organism:{o.id}", " · ".join(filter(None, [o.code or f"{o.count} {i18n.translate_value(place.nouns)}", o.sex, o.genotype]))))
        codes = {h.id: h.code for h in session.scalars(select(OrgHousing).where(OrgHousing.id.in_(list(houses) or [0])))}
        groups = [(str(hid), f"{i18n.translate_value(place.housing).capitalize()} {codes.get(hid, hid)}", len(v)) for hid, v in sorted(houses.items(), key=lambda kv: positions.natural(codes.get(kv[0], kv[0])))]
        cohort_noun = i18n.translate_value((place.mv.cohort_noun if place.mv else "") or "cohort").capitalize()
        for c in positions.natural_sorted(session.scalars(select(OrgCohort).where(OrgCohort.module_id_fk == place.module.id)),
                                          key=lambda c: c.code):
            if f"cohort:{c.id}" not in have and c.count_current:
                single.append((f"cohort:{c.id}", f"{cohort_noun} {c.code} · {c.count_current}"))
    return {"groups": groups, "single": single}


# ---------------------------------------------------------------- readings

def readings(session, exp: Experiment, place: Place, subs: list[Subject] | None = None,
             readout: dict | None = None) -> dict[str, dict[date, float]]:
    """{subject key: {date: value}} for the readout (or another measure)."""
    subs = subjects(session, exp, place) if subs is None else subs
    out: dict[str, dict[date, float]] = {s.key: {} for s in subs}
    readout = readout or readout_of(exp, place)
    if place.subject_kind == "mouse" and readout["key"] == "body_weight":
        ids = {s.id: s.key for s in subs}
        if ids:
            for w in session.scalars(select(MouseWeight).where(MouseWeight.mouse_id_fk.in_(list(ids)))):
                out[ids[w.mouse_id_fk]][w.weigh_date] = w.grams
        return out
    rows = list(session.scalars(select(ExperimentReading).where(ExperimentReading.experiment_id_fk == exp.id,
                                                                  ExperimentReading.readout_key == readout["key"])))
    if readout.get("per_group"):
        # Against the control group's mean ΔCt that day: worked out afresh,
        # so a change of group or control is followed (a day with no
        # control has nothing to show).
        groups = {x.key: x.group for x in subs}
        by_day: dict[date, dict[str, dict]] = {}
        for r in rows:
            inputs = derived.read_note(r.note)
            if r.subject in out and inputs is not None:
                by_day.setdefault(r.read_on, {})[r.subject] = inputs
        for on, day in by_day.items():
            for key, value in (_group_values(readout, day, groups) or {}).items():
                out[key][on] = value
        return out
    for r in rows:
        if r.subject in out:
            out[r.subject][r.read_on] = r.value
    return out


def reading_inputs(session, exp: Experiment, place: Place, subs: list[Subject] | None = None,
                   readout: dict | None = None) -> dict[str, dict[date, dict]]:
    """{subject key: {date: inputs}} for a readout BioManager works out:
    what was measured, kept with each reading."""
    readout = readout or readout_of(exp, place)
    if not readout.get("formula"):
        return {}
    keys = {x.key for x in (subjects(session, exp, place) if subs is None else subs)}
    out: dict[str, dict[date, dict]] = {}
    for r in session.scalars(select(ExperimentReading).where(ExperimentReading.experiment_id_fk == exp.id,
                                                             ExperimentReading.readout_key == readout["key"])):
        inputs = derived.read_note(r.note)
        if r.subject in keys and inputs is not None:
            out.setdefault(r.subject, {})[r.read_on] = inputs
    return out


def _group_values(readout: dict, day: dict[str, dict], groups: dict[str, str]) -> dict[str, float] | None:
    """ΔΔCt or fold change for each animal of one day, from {subject: inputs};
    None when the day has no control (or no control is set)."""
    keys = list(day)
    rows = [{"group": groups.get(k, ""), **day[k]} for k in keys]
    try:
        return dict(zip(keys, derived.group_values(readout["formula"], rows, readout.get("control") or "")))
    except derived.DerivedError:
        return None


def is_control(readout: dict, subject: Subject) -> bool:
    return bool(readout.get("control")) and subject.group.strip().lower() == readout["control"].strip().lower()


def control_first(readout: dict, items: list, subject_of) -> list:
    """Readings to save in an order that works: for a readout against the
    control group, its animals first (the others are worked out from them)."""
    if not readout.get("per_group"):
        return items

    def rank(item) -> int:
        x = subject_of(item)
        return 0 if x is not None and is_control(readout, x) else 1
    return sorted(items, key=rank)


def weight_on(values: dict[date, float], day: date | None) -> tuple[float | None, date | None]:
    """The latest value on or before `day` (the latest of all without one)."""
    usable = sorted(d for d in values if day is None or d <= day)
    return (values[usable[-1]], usable[-1]) if usable else (None, None)


def set_reading(session, exp: Experiment, place: Place, subject: Subject, on: date, raw,
                readout: dict | None = None) -> str | None:
    """Store (or, blank, remove) one day's readout (or, given `readout`,
    another measure's); a sentence when it can't. `raw` is the value typed,
    or for a readout BioManager works out, what was measured
    ({"length": 12.1, "width": 8.4})."""
    readout = readout or readout_of(exp, place)
    if readout.get("formula"):
        return _set_derived(session, exp, place, readout, subject, on, raw)
    if isinstance(raw, dict):
        return gettext("%(readout)s is typed as one number for each, not worked out from measurements.",
                       readout=i18n.translate_value(readout["label"]))
    raw = ("" if raw is None else str(raw)).strip().replace(",", ".")
    value = None
    if raw:
        try:
            value = float(raw)
        except ValueError:
            return gettext("“%(value)s” isn't a number.", value=raw)
        if value < 0:
            return gettext("A readout can't be below zero.")
        if readout["kind"] == "fraction" and subject.start is not None and value > subject.start:
            return gettext("%(label)s started with %(start)s: %(value)s is more than that.",
                           label=subject.label, start=subject.start, value=f"{value:g}")
        if readout["key"] == "body_weight" and place.subject_kind == "mouse" and not 0 < value < 200:
            return gettext("%(value)s g isn't a mouse's weight.", value=f"{value:g}")
    if place.subject_kind == "mouse" and readout["key"] == "body_weight":
        if not subject.can_edit:
            return NotYours(f"{subject.label}: {access.reason_denied(subject.record)}")
        row = session.scalar(select(MouseWeight).where(MouseWeight.mouse_id_fk == subject.id, MouseWeight.weigh_date == on))
        if value is None:
            if row is not None:
                session.delete(row)
        elif row is None:
            session.add(MouseWeight(mouse_id_fk=subject.id, weigh_date=on, grams=value,
                                    notes=f"{exp.name}"[:200], recorded_by=g.user.username))
        else:
            row.grams, row.recorded_by = value, g.user.username
        return None
    row = session.scalar(select(ExperimentReading).where(
        ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.subject == subject.key,
        ExperimentReading.readout_key == readout["key"], ExperimentReading.read_on == on))
    if value is None:
        if row is not None:
            session.delete(row)
    elif row is None:
        session.add(ExperimentReading(experiment_id_fk=exp.id, subject=subject.key, readout_key=readout["key"],
                                      read_on=on, value=value, recorded_by=g.user.username))
    else:
        row.value, row.recorded_by = value, g.user.username
    return None


def _blank(raw) -> bool:
    if raw is None:
        return True
    if isinstance(raw, dict):
        return all(v is None or not str(v).strip() for v in raw.values())
    return not str(raw).strip()


def derive_day(session, exp: Experiment, place: Place, on: date, given: dict[str, dict],
               subs: list[Subject] | None = None, readout: dict | None = None) -> dict[str, float]:
    """What a day's readings would be, for a readout (or another measure)
    BioManager works out: {subject key: value} for the animals `given`
    ({key: inputs}, checked), against what that day already has (a control
    group's mean ΔCt takes every control animal of the day). A
    DerivedError, in words, otherwise."""
    readout = readout or readout_of(exp, place)
    f = derived.formula(readout["formula"])
    if not f.per_group:
        return {k: derived.compute(f.key, v) for k, v in given.items()}
    if not readout.get("control") and readout.get("measure"):
        raise derived.DerivedError(gettext("Say which group is the control of %(measure)s (its Control group): it is worked out against that group's mean ΔCt.",  # noqa: E501
                                           measure=readout["label"]))
    if not readout.get("control"):
        raise derived.DerivedError(gettext("Say which group is the control (Control group, beside the readout): %(readout)s is worked out against its mean ΔCt.",  # noqa: E501
                                           readout=i18n.translate_value(f.label)))
    subs = subjects(session, exp, place) if subs is None else subs
    groups = {x.key: x.group for x in subs}
    day = {k: v[on] for k, v in reading_inputs(session, exp, place, subs, readout).items() if on in v}
    day.update(given)
    if not any(groups.get(k, "").strip().lower() == readout["control"].strip().lower() for k in day):
        raise derived.DerivedError(gettext("No animal of the control group “%(group)s” has its Ct values on %(day)s yet: enter theirs first.",  # noqa: E501
                                           group=readout["control"], day=on.isoformat()))
    values = _group_values(readout, day, groups) or {}
    return {k: values[k] for k in given if k in values}


_TWO_CTS = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*(?:[-−–/;:]|\s)\s*(\d+(?:[.,]\d+)?)\s*$")


def typed_cts(f: derived.Formula, raw) -> dict | float:
    """What was typed in a Ct measure's cell: two Cts ("24.8 − 18.2", the
    target's then the reference's) as inputs, or for ΔCt the ΔCt itself
    as a number. A DerivedError, in words, for anything else."""
    text = str(raw or "").strip().replace("−", "-")
    two = _TWO_CTS.match(text)
    if two:
        return {"ct_target": two.group(1).replace(",", "."), "ct_reference": two.group(2).replace(",", ".")}
    try:
        value = float(text.replace(",", "."))
    except ValueError:
        value = None
    if f.key == "delta_ct" and value is not None and math.isfinite(value):
        if abs(value) > derived.CT_MAX:
            raise derived.DerivedError(gettext("A ΔCt of %(value)s isn't one a run gives (at most %(max)s cycles either way).",
                                               value=f"{value:g}", max=f"{derived.CT_MAX:g}"))
        return derived.tidy(value)
    raise derived.DerivedError(gettext("Type the two Cts, the target's then the reference's, e.g. 24.8 − 18.2: %(readout)s is worked out from them.",  # noqa: E501
                                       readout=i18n.translate_value(f.label)))


def _set_derived(session, exp: Experiment, place: Place, readout: dict, subject: Subject, on: date, raw) -> str | None:
    """A reading of a readout BioManager works out: the inputs checked, the
    value computed and kept with them (the note, as JSON); a day against
    the control group is worked out again whole."""
    f = derived.formula(readout["formula"])
    row = session.scalar(select(ExperimentReading).where(
        ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.subject == subject.key,
        ExperimentReading.readout_key == readout["key"], ExperimentReading.read_on == on))
    subs = subjects(session, exp, place) if f.per_group else None
    if _blank(raw):
        if row is None:
            return None
        if f.per_group and is_control(readout, subject):
            day = {k: v[on] for k, v in reading_inputs(session, exp, place, subs, readout).items() if on in v}
            controls = [x for x in subs if x.key in day and is_control(readout, x)]
            if len(controls) == 1 and len(day) > 1:
                return gettext("The other animals' %(readout)s on %(day)s is worked out against the control group: remove theirs first, or keep a control.",  # noqa: E501
                               readout=i18n.translate_value(f.label), day=on.isoformat())
        session.delete(row)
        if f.per_group:
            session.flush()
            _store_day(session, exp, readout, on, subs)
        return None
    try:
        if readout.get("measure") and not isinstance(raw, dict):
            raw = typed_cts(f, raw)
        if isinstance(raw, float):            # a ΔCt typed as it is: nothing to work it out from
            inputs, value = None, raw
        else:
            inputs = derived.check(f.key, raw)
            value = derive_day(session, exp, place, on, {subject.key: inputs}, subs, readout)[subject.key]
    except derived.DerivedError as e:
        return f"{subject.label}: {e}"
    if value < 0 and not f.can_be_negative:
        return gettext("A readout can't be below zero.")
    note = derived.note(inputs) if inputs is not None else ""
    if row is None:
        session.add(ExperimentReading(experiment_id_fk=exp.id, subject=subject.key, readout_key=readout["key"],
                                      read_on=on, value=value, note=note, recorded_by=g.user.username))
    else:
        row.value, row.note, row.recorded_by = value, note, g.user.username
    if f.per_group:
        session.flush()
        _store_day(session, exp, readout, on, subs)
    return None


def _store_day(session, exp: Experiment, readout: dict, on: date, subs: list[Subject]) -> None:
    """Keep a day's ΔΔCt or fold change as worked out now, for every animal
    of it (a new control animal moves the mean the others are against)."""
    groups = {x.key: x.group for x in subs}
    rows = {r.subject: r for r in session.scalars(select(ExperimentReading).where(
        ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.readout_key == readout["key"],
        ExperimentReading.read_on == on))}
    day = {}
    for key, r in rows.items():
        inputs = derived.read_note(r.note)
        if inputs is not None:
            day[key] = inputs
    for key, value in (_group_values(readout, day, groups) or {}).items():
        rows[key].value = value


def readout_table(session, exp: Experiment, place: Place, subs: list[Subject] | None = None,
                  readout: dict | None = None) -> dict:
    """The readout (or another measure) by animal and date, each date's day
    number, and each value as a percentage: of the animal's first value, or
    for a count of how many there were at the start."""
    subs = subjects(session, exp, place) if subs is None else subs
    readout = readout or readout_of(exp, place)
    series = readings(session, exp, place, subs, readout)
    measured = reading_inputs(session, exp, place, subs, readout)
    dates = sorted({d for values in series.values() for d in values})
    start = exp.start_date
    rows = []
    for s in subs:
        values = [series[s.key].get(d) for d in dates]
        if readout["kind"] == "fraction":
            pct = [round(v / s.start * 100, 1) if (v is not None and s.start) else None for v in values]
        else:
            first = next((v for v in values if v is not None), None)
            pct = [round(v / first * 100, 1) if (v is not None and first) else None for v in values]
        rows.append({"key": s.key, "mouse": s.id if s.kind == "mouse" else None,
                     "mouse_id": s.label.lstrip("#") if s.kind == "mouse" else s.label, "label": s.label,
                     "group": s.group, "sex": s.sex, "start": s.start, "values": values, "pct": pct,
                     "left_out": s.left_out,
                     **({"inputs": [measured.get(s.key, {}).get(d) for d in dates]} if readout.get("formula") else {})})
    table = {"dates": [d.isoformat() for d in dates],
             "days": [((d - start).days + 1) if start else None for d in dates], "rows": rows, "readout": readout}
    from . import exp_stats
    table["stats"] = {"raw": exp_stats.by_day(table), "pct": exp_stats.by_day(table, percent=True)}
    return table


# ---------------------------------------------------------------- the page's data

def payload(session, exp: Experiment, place: Place) -> dict:
    from . import experiment_plan
    from . import experiment_steps as xs
    subs = subjects(session, exp, place)
    readout = readout_of(exp, place)
    groups = sorted({s.group for s in subs if s.group})
    table = readout_table(session, exp, place, subs)
    return {
        "id": exp.id, "name": exp.name, "status": exp.status, "owner": exp.owner_username,
        "db": place.key, "db_label": place.label, "family": place.family, "noun": place.noun, "nouns": place.nouns,
        "housing": place.housing, "editable": access.can_edit_experiment(exp),
        "start_date": exp.start_date.isoformat() if exp.start_date else "",
        "readout": readout, "weighs": readout["key"] == "body_weight",
        "subjects": [s.as_dict() for s in subs], "groups": groups,
        "kinds": [{"key": k, "label": label, "icon": icon} for k, label, icon in kinds_for(place.family)],
        "steps": [xs.step_dict(st, place.family, session) for st in exp.steps],
        "schedule": xs.schedule(session, exp, place, subs),
        "table": table,
        "measures": [{**m, "table": readout_table(session, exp, place, subs, m),
                      "count": measure_count(session, exp, m["key"])} for m in measures_of(exp)],
        "regimens": xs.regimens_for(session, place.family),
        "reagents": reagent_choices(session),
        "sample_inventories": sample_inventories(session),
        **experiment_plan.page_data(session, exp, place, subs, table),
    }


def measure_count(session, exp: Experiment, key: str) -> int:
    """How many readings a measure holds (what deleting it takes with it)."""
    return session.scalar(select(func.count(ExperimentReading.id)).where(
        ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.readout_key == key)) or 0


def reagent_choices(session, limit: int = 400) -> list[dict]:
    """Inventory items to say a manipulation used: the lab's reagents,
    antibodies and the like, with their lot, newest first."""
    from . import inventory_service as inv
    from .models import InventoryItem
    out = []
    for module in inv.list_modules(session):
        if module.kind in ("orders", "samples") or not lab.can_see(module):
            continue
        mv = inv.view(module)
        rows = session.scalars(select(InventoryItem).where(InventoryItem.module_id_fk == module.id)
                               .order_by(InventoryItem.id.desc()).limit(limit)).all()
        for item in rows:
            if mv.is_available(item.status) is False:
                continue
            out.append({"id": item.id, "label": " · ".join(filter(None, [item.name, f"lot {item.lot}" if item.lot else "",
                                                                         mv.label])),
                        "name": item.name or "", "lot": item.lot or "",
                        "expires": item.expires_on.isoformat() if item.expires_on else ""})
    return out[:limit]


def sample_inventories(session) -> list[dict]:
    """Inventories a sample record can go in: those with a Source column."""
    from . import inventory_service as inv
    out = []
    for module in inv.list_modules(session):
        if not lab.can_see(module):
            continue
        mv = inv.view(module)
        if any(f["type"] == "source" for f in mv.fields):
            out.append({"key": module.key, "label": mv.label})
    return out


def _load(session, experiment_id: int) -> tuple[Experiment, Place]:
    exp = session.get(Experiment, experiment_id)
    if exp is None:
        abort(404)
    place = place_for(session, exp.db or "colony")
    if place is None:
        abort(404)
    return exp, place


def _refuse(exp: Experiment):
    if access.can_edit_experiment(exp):
        return None
    return jsonify({"ok": False, "error": access.denied_message("experiment", exp.owner_username)}), 403


@bp.before_request
def require_login():
    if g.get("user") is None:
        return redirect(url_for("login", next=request.path))
    return None


# ---------------------------------------------------------------- pages

def render_page(experiment_id: int):
    """The experiment page, for every database (the colony's route calls this)."""
    from . import experiment_steps as xs
    with SessionLocal() as s:
        exp = s.get(Experiment, experiment_id)
        if exp is None:
            flash(gettext("Experiment not found."), "error")
            return redirect(url_for("colony", view="experiments"))
        place = place_for(s, exp.db or "colony")
        if place is None:
            abort(404)
        data = payload(s, exp, place)
        info = {"id": exp.id, "name": exp.name, "description": exp.description, "treatment_plan": exp.treatment_plan,
                "status": exp.status, "owner": exp.owner_username, "editable": data["editable"],
                "start_date": data["start_date"], "end_date": exp.end_date.isoformat() if exp.end_date else ""}
        statuses = ["planned", "cancelled"] if exp.status == "planned" else STARTED
        return render_template(
            "experiment.html", experiment=info, place=place, data=data, statuses=statuses,
            readouts=readout_choices(place), derived_readouts=derived_choices(place), readout=data["readout"], add=candidates(s, exp, place) if data["editable"] else None,
            notebook_on=lab.feature_on(s, "notebook"), notebook_page=xs.my_notebook_page(s, exp.id),
            back_url=place.list_url)


@bp.get("/<int:experiment_id>")
def page(experiment_id: int):
    with SessionLocal() as s:
        exp = s.get(Experiment, experiment_id)
        if exp is not None and (exp.db or "colony") == "colony":
            return redirect(url_for("experiment_detail", experiment_id=experiment_id))
    return render_page(experiment_id)


@bp.get("/<int:experiment_id>/data.json")
def data_json(experiment_id: int):
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        return jsonify({"ok": True, **payload(s, exp, place)})


def tab_context(session, place: Place) -> dict:
    """What a database's Experiments tab shows (templates/_experiments_tab.html):
    its experiments, newest first by status, and the readouts a new one can
    follow. The zebrafish, fly and worm, and organism pages pass it on as
    `experiments_tab`; the mouse colony has its own tab."""
    rows = session.scalars(select(Experiment).where(Experiment.db == place.key)
                           .order_by(Experiment.status, Experiment.created_at.desc())).all()
    counts = dict(session.execute(
        select(ExperimentSubject.experiment_id_fk, func.count(ExperimentSubject.id))
        .where(ExperimentSubject.experiment_id_fk.in_([e.id for e in rows]))
        .group_by(ExperimentSubject.experiment_id_fk)).all()) if rows else {}
    cards = [{"id": e.id, "name": e.name, "description": e.description, "status": e.status,
              "owner": e.owner_username, "editable": access.can_edit_experiment(e), "count": counts.get(e.id, 0),
              "start_date": i18n.strftime(e.start_date, "%b %d, %Y") if e.start_date else "",
              "readout": readout_of(e, place)["label"], "url": url_for("experiments.page", experiment_id=e.id)}
             for e in rows]
    # A new experiment can start with a whole tank, rack or housing of them,
    # as a mouse one starts with a cage's mice.
    groups = candidates(session, Experiment(id=0, db=place.key), place)["groups"]
    return {"place": place, "cards": cards, "readouts": readout_choices(place), "groups": groups,
            "today": date.today().isoformat(), "statuses": STATUSES}


def tab_count(session, db_key: str) -> int:
    """How many experiments a database has, for the count on its tab."""
    return session.scalar(select(func.count(Experiment.id)).where(Experiment.db == db_key)) or 0


@bp.get("/in/<db_key>")
def index(db_key: str):
    """A database's experiments are on its Experiments tab; this address
    (from before the tab, in bookmarks and notebook links) goes there."""
    if db_key == "colony":
        return redirect(url_for("colony", view="experiments"))
    with SessionLocal() as s:
        place = place_for(s, db_key)
        if place is None:
            abort(404)
        return redirect(place.list_url)


@bp.post("/in/<db_key>/create")
def create(db_key: str):
    with SessionLocal() as s:
        place = place_for(s, db_key)
        if place is None:
            abort(404)
        form = request.form
        from .services import parse_date
        readout = form.get("readout") or ""
        exp = Experiment(name=(form.get("name") or "").strip()[:200] or "Untitled experiment",
                         description=(form.get("description") or "").strip(),
                         treatment_plan=(form.get("treatment_plan") or "").strip(), status="active",
                         owner_username=g.user.username, start_date=parse_date(form.get("start_date")),
                         db=place.key, readout=readout_json(readout, form.get("readout_control") or "") or "")
        s.add(exp)
        s.flush()
        start_with = (form.get("from_group") or "").strip()
        if start_with:
            for kind, rec, start in _records_to_add(s, place, "group", start_with):
                s.add(ExperimentSubject(experiment_id_fk=exp.id, subject_kind=kind, subject_id=rec.id,
                                        start_count=start))
        s.commit()
        return redirect(page_url(exp))


@bp.post("/in/<db_key>/design")
def design(db_key: str):
    """A whole design as one planned experiment (an assistant's
    `experiment_new`): JSON {name, description, notes, via, readout
    ({"key"} or {"label", "unit", "kind"}), readout_day, planned_test,
    start_date, groups: [{"name", "add": [{"how": "one" | "group", "value"}]}],
    steps: [{days, kind, agent, dose, route, concentration, group, notes}]}.
    The animals go in as Add does and the steps as the Regimen saves them,
    so their rules hold; nothing is due until a person presses Start."""
    from . import experiment_plan as plan
    from . import experiment_steps as xs
    from .services import parse_date
    data = request.get_json(silent=True) or {}
    with SessionLocal() as s:
        place = place_for(s, db_key)
        if place is None:
            return jsonify({"ok": False, "error": gettext("There is no database “%(key)s” to plan an experiment in.", key=db_key)}), 404
        name = str(data.get("name") or "").strip()[:200]
        if not name:
            return jsonify({"ok": False, "error": gettext("Give the experiment a name.")}), 400
        readout = data.get("readout") or {}
        if isinstance(readout, str):
            readout = {"key": readout}
        if not isinstance(readout, dict):
            readout = {}
        # A built-in readout, one BioManager works out ("derived:tumour_lw", with its control), or your own.
        chosen_json = readout_json(str(readout.get("key") or ""), str(readout.get("control") or ""))
        if chosen_json is None and readout.get("label"):
            kind = readout.get("kind") if readout.get("kind") in ("value", "fraction") else "value"
            chosen_json = json.dumps({"key": "custom", "label": str(readout["label"])[:60],
                                      "unit": str(readout.get("unit") or "")[:20], "kind": kind})
        elif chosen_json is None and readout.get("key"):
            return jsonify({"ok": False, "error": gettext("“%(readout)s” is not a readout; give its label, unit and kind for your own.", readout=readout.get("key"))}), 400  # noqa: E501
        chosen_json = chosen_json or ""
        test = str(data.get("planned_test") or "").strip()
        if test and test not in plan.TESTS:
            return jsonify({"ok": False, "error": gettext("“%(test)s” is not a test the app runs.", test=test)}), 400
        day = data.get("readout_day")
        if day not in (None, "") and not (str(day).isdigit() and 0 < int(day) <= 1000):
            return jsonify({"ok": False, "error": gettext("“%(value)s” isn't a day of the experiment (day 1 is the start date).", value=day)}), 400  # noqa: E501
        description = str(data.get("description") or "").strip()
        via = str(data.get("via") or "").strip()[:80]
        if via:
            description = "\n\n".join(filter(None, [description, f"Designed with {via}."]))
        exp = Experiment(name=name, description=description, treatment_plan=str(data.get("notes") or "").strip(),
                         status="planned", owner_username=g.user.username,
                         start_date=parse_date(str(data.get("start_date") or "")), db=place.key, readout=chosen_json,
                         readout_day=int(day) if day not in (None, "") else None, planned_test=test)
        s.add(exp)
        s.flush()
        said = []
        for group in [x for x in data.get("groups") or [] if isinstance(x, dict)]:
            gname = str(group.get("name") or "").strip()[:80]
            for add in [x for x in group.get("add") or [] if isinstance(x, dict)]:
                added, skipped = add_animals(s, exp, place, str(add.get("how") or "one"), str(add.get("value") or ""), gname)
                if skipped:
                    said.append(gettext("%(count)s skipped: not yours to change.", count=skipped))
                if not added and not skipped:
                    said.append(gettext("Nothing to add from %(value)s for %(group)s.", value=add.get("value"), group=gname or "-"))
                s.flush()
        for st in [x for x in data.get("steps") or [] if isinstance(x, dict)]:
            try:
                days = xs.parse_days(str(st.get("days", "")))
            except ValueError as error:
                s.rollback()
                return jsonify({"ok": False, "error": str(error)}), 400
            s.refresh(exp)
            _step, error = xs._fill_step(s, exp, place, {**st, "id": 0}, days)
            if error:
                s.rollback()
                return error
        s.commit()
        return jsonify({"ok": True, "id": exp.id, "url": page_url(exp), "warnings": said})


# ---------------------------------------------------------------- changing it

def status_refusal(exp: Experiment, status: str) -> str | None:
    """A planned experiment starts only with Start (a person's button), and
    one that has started is never planned again."""
    if status == (exp.status or "") or not status:
        return None
    if exp.status == "planned" and status != "cancelled":
        return gettext("“%(name)s” is planned: press Start on its page to start it.", name=exp.name)
    if status == "planned":
        return gettext("“%(name)s” has started, so it can't be planned again.", name=exp.name)
    return None


@bp.post("/<int:experiment_id>/update")
def update(experiment_id: int):
    """The details, autosaved one field at a time, and the readout. After
    Start, a change to the main readout, its day or the test planned asks
    for a reason (`reason`; app/experiment_plan.py)."""
    from . import experiment_plan as plan
    from .services import parse_date
    form = request.form
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        start = parse_date(form.get("start_date")) if "start_date" in form else exp.start_date
        end = parse_date(form.get("end_date")) if "end_date" in form else exp.end_date
        status = (form.get("status") or "").strip().lower() if "status" in form else exp.status
        if start and end and end < start:
            return jsonify({"ok": False, "error": gettext("The end date (%(end)s) is before the start date (%(start)s).",
                                                          end=end.isoformat(), start=start.isoformat())}), 409
        if status and status not in STATUSES:
            return jsonify({"ok": False, "error": gettext("“%(status)s” is not an experiment status.", status=status)}), 409
        problem = status_refusal(exp, status)
        if problem:
            return jsonify({"ok": False, "error": problem}), 409
        readout_day = exp.readout_day
        if "readout_day" in form:
            raw = (form.get("readout_day") or "").strip().lower().removeprefix("day").strip()
            if raw and not (raw.lstrip("-").isdigit() and 0 < int(raw) <= 1000):
                return jsonify({"ok": False, "error": gettext("“%(value)s” isn't a day of the experiment (day 1 is the start date).", value=raw)}), 409  # noqa: E501
            readout_day = int(raw) if raw else None
        planned_test = exp.planned_test or ""
        if "planned_test" in form:
            planned_test = (form.get("planned_test") or "").strip()
            if planned_test and planned_test not in plan.TESTS:
                return jsonify({"ok": False, "error": gettext("“%(test)s” is not a test the app runs.", test=planned_test)}), 409
        was = {"readout": plan.readout_text(readout_of(exp, place)), "readout_day": str(exp.readout_day or ""),
               "test": plan.test_name(exp.planned_test or "")}
        for name in ("name", "description", "treatment_plan"):
            if name in form:
                value = (form.get(name) or "").strip() if name == "name" else form.get(name, "")
                if name != "name" or value:
                    setattr(exp, name, value[:200] if name == "name" else value)
        if "readout" in form:
            key = form.get("readout") or ""
            before = exp.readout
            if key == "custom":
                kind = form.get("readout_kind") if form.get("readout_kind") in ("value", "fraction") else "value"
                exp.readout = json.dumps({"key": "custom", "label": (form.get("readout_label") or "Measurement")[:60],
                                          "unit": (form.get("readout_unit") or "")[:20], "kind": kind})
            elif readout_json(key) is not None:
                exp.readout = readout_json(key, form.get("readout_control") or "")
            if exp.readout != before:
                exp.readout = keep_measures(exp.readout, before)
        exp.readout_day, exp.planned_test = readout_day, planned_test
        exp.status, exp.start_date, exp.end_date = status or "active", start, end
        exp.updated_at = datetime.utcnow()
        now = {"readout": plan.readout_text(readout_of(exp, place)), "readout_day": str(exp.readout_day or ""),
               "test": plan.test_name(exp.planned_test or "")}
        for what in ("readout", "readout_day", "test"):
            problem = plan.note(s, exp, what, was[what], now[what], form.get("reason", ""))
            if problem:
                s.rollback()
                return plan.ask_response(problem)
        s.commit()
        return jsonify({"ok": True, **payload(s, exp, place)})


def add_animals(session, exp: Experiment, place: Place, how: str, value: str, group: str) -> tuple[int, int]:
    """Put a cage, tank, rack or housing of animals ("group"), or one
    ("one"), in the experiment: (added, skipped as not this person's)."""
    from .app import can_edit_mouse, mouse_is_active
    session.flush()
    session.expire(exp, ["memberships"])        # mice added a moment ago count as in it
    have = {x.key for x in subjects(session, exp, place)}
    added = skipped = 0
    if place.subject_kind == "mouse":
        if how == "group":
            cage = session.scalar(select(CageRecord).where(CageRecord.cage_id == value))
            mice = list(cage.mice) if cage else []
        else:
            mice = [m for m in [session.get(MouseRecord, int(value))] if m] if value.isdigit() else []
        for m in mice:
            if f"mouse:{m.id}" in have or not mouse_is_active(m):
                continue
            if not can_edit_mouse(m):
                skipped += 1
                continue
            session.add(ExperimentMouse(experiment_id_fk=exp.id, mouse_id_fk=m.id, treatment_group=group))
            have.add(f"mouse:{m.id}")
            added += 1
    else:
        for kind, rec, start in _records_to_add(session, place, how, value):
            key = f"{kind}:{rec.id}"
            if key in have:
                continue
            session.add(ExperimentSubject(experiment_id_fk=exp.id, subject_kind=kind, subject_id=rec.id,
                                          treatment_group=group, start_count=start))
            have.add(key)
            added += 1
    return added, skipped


@bp.post("/<int:experiment_id>/subjects/add")
def add_subjects(experiment_id: int):
    """A cage, tank, rack or housing of animals ("group"), or one ("one").
    After Start, with a plan, it asks for a reason (`reason`)."""
    from . import experiment_plan as plan
    data = request.get_json(silent=True) or request.form.to_dict()
    how, value, group = data.get("how"), str(data.get("value") or ""), str(data.get("group") or "").strip()[:80]
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        before = plan.members(s, exp, place) if plan.fixed(exp) else {}
        added, skipped = add_animals(s, exp, place, how, value, group)
        problem = plan.groups_changed(s, exp, place, before, str(data.get("reason") or "")) if added else None
        if problem:
            s.rollback()
            return plan.ask_response(problem)
        s.commit()
        s.refresh(exp)
        if added:
            message = gettext("Added %(count)s %(nouns)s.", count=added,
                              nouns=i18n.translate_value(place.noun if added == 1 else place.nouns))
        else:
            message = gettext("No %(nouns)s to add there.", nouns=i18n.translate_value(place.nouns))
        if skipped:
            message += " " + gettext("%(count)s skipped: not yours to change.", count=skipped)
        return jsonify({"ok": True, "message": message, **payload(s, exp, place),
                        "add": candidates(s, exp, place)})


def _records_to_add(session, place: Place, how: str, value: str) -> list[tuple[str, object, int | None]]:
    """(kind, record, animals at the start) for the fish rows of a tank, the
    vials in a rack, the animals in a housing, or one ("fish:12",
    "clutch:3", "organism:40", "cohort:2", or a bare id)."""
    kind, _, raw = value.rpartition(":")
    if not raw.isdigit():
        return []
    ident = int(raw)
    if how != "group" and kind == "clutch" and place.subject_kind == "fish":
        c = session.get(ClutchRecord, ident)
        return [("clutch", c, c.larvae_count or c.embryo_count or None)] if c else []
    if how != "group" and kind == "cohort" and place.subject_kind == "organism":
        c = session.get(OrgCohort, ident)
        return [("cohort", c, c.count_initial or c.count_current or None)] if c and c.module_id_fk == place.module.id else []
    return [(place.subject_kind, rec, start) for rec, start in _plain_records(session, place, how, ident)]


def _plain_records(session, place: Place, how: str, ident: int) -> list[tuple[object, int | None]]:
    if place.subject_kind == "fish":
        if how == "group":
            tank = session.get(TankRecord, ident)
            return [(f, f.count) for f in (tank.fish if tank else []) if f.status == "alive"]
        f = session.get(FishRecord, ident)
        return [(f, f.count)] if f else []
    if place.subject_kind == "unit":
        query = select(StockUnit).where(StockUnit.module_id_fk == place.module.id, StockUnit.active.is_(True))
        query = query.where(StockUnit.rack_id_fk == ident) if how == "group" else query.where(StockUnit.id == ident)
        return [(u, None) for u in session.scalars(query)]
    query = select(Organism).where(Organism.module_id_fk == place.module.id)
    query = query.where(Organism.housing_id_fk == ident) if how == "group" else query.where(Organism.id == ident)
    return [(o, o.count or 1) for o in session.scalars(query) if place.mv.is_alive(o)]


@bp.post("/<int:experiment_id>/subjects/<path:key>/update")
def update_subject(experiment_id: int, key: str):
    """Its group (`whole_group`: every one in its group), note, or count at
    the start. After Start, with a plan, a group or count asks for a reason."""
    from . import experiment_plan as plan
    data = request.get_json(silent=True) or request.form.to_dict()
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        subject = find_subject(s, exp, place, key)
        if subject is None:
            return jsonify({"ok": False, "error": gettext("That one is no longer in the experiment.")}), 404
        before = plan.members(s, exp, place) if plan.fixed(exp) else {}
        row = subject.row
        if "group" in data:
            new_group = str(data.get("group") or "").strip()[:80]
            # The whole group at once: renaming it, or merging it into another.
            if data.get("whole_group") and subject.group.strip():
                for other in subjects(s, exp, place, group=subject.group):
                    other.row.treatment_group = new_group
            row.treatment_group = new_group
        if "note" in data:
            row.note = str(data.get("note") or "").strip()[:200]
        if "start" in data and isinstance(row, ExperimentSubject):
            raw = str(data.get("start") or "").strip()
            if raw and not raw.isdigit():
                return jsonify({"ok": False, "error": gettext("“%(value)s” isn't a number of animals.", value=raw)}), 400
            row.start_count = int(raw) if raw else None
        problem = plan.groups_changed(s, exp, place, before, str(data.get("reason") or ""))
        if problem:
            s.rollback()
            return plan.ask_response(problem)
        s.commit()
        s.refresh(exp)
        return jsonify({"ok": True, **payload(s, exp, place)})


@bp.post("/<int:experiment_id>/subjects/<path:key>/remove")
def remove_subject(experiment_id: int, key: str):
    from . import experiment_plan as plan
    data = request.get_json(silent=True) or request.form.to_dict()
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        subject = find_subject(s, exp, place, key)
        if subject is not None:
            before = plan.members(s, exp, place) if plan.fixed(exp) else {}
            s.delete(subject.row)
            problem = plan.groups_changed(s, exp, place, before, str(data.get("reason") or ""))
            if problem:
                s.rollback()
                return plan.ask_response(problem)
            s.commit()
            s.refresh(exp)
        return jsonify({"ok": True, **payload(s, exp, place), "add": candidates(s, exp, place)})


@bp.post("/<int:experiment_id>/subjects/<path:key>/leave-out")
def leave_out(experiment_id: int, key: str):
    """{"reason": why, "back": false}: leave an animal out of the analysis
    (the tests by day and the chart; its readout stays), or put it back.
    Leaving out always says why; after Start, with a plan, it is a change
    from the plan, and so is putting it back."""
    from . import experiment_plan as plan
    data = request.get_json(silent=True) or request.form.to_dict()
    reason = str(data.get("reason") or "").strip()[:200]
    back = str(data.get("back") or "").lower() in ("1", "true", "yes", "on")
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        subject = find_subject(s, exp, place, key)
        if subject is None:
            return jsonify({"ok": False, "error": gettext("That one is no longer in the experiment.")}), 404
        if not back and not reason:
            return jsonify({"ok": False, "needs_reason": True,
                            "error": gettext("Say why %(label)s is left out of the analysis.", label=subject.label)}), 409
        was = plan.LEFT_OUT if subject.left_out else plan.IN_ANALYSIS
        subject.row.left_out = "" if back else reason
        now = plan.IN_ANALYSIS if back else plan.LEFT_OUT
        line = subject.label if back else f"{subject.label}: {reason}"
        problem = plan.note(s, exp, "left_out", was, now, reason, [line]) if was != now else None
        if problem:
            s.rollback()
            return plan.ask_response(problem)
        exp.updated_at = datetime.utcnow()
        s.commit()
        s.refresh(exp)
        return jsonify({"ok": True, **payload(s, exp, place)})


@bp.post("/<int:experiment_id>/start")
def start(experiment_id: int):
    """Start a planned experiment: day 1 is today unless a start date is
    set; it becomes active, and a copy of its design is kept, so a change
    from it later asks why (app/experiment_plan.py). Only from the page:
    no assistant's change does this."""
    from . import experiment_plan as plan
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        if not access.can_edit_experiment(exp):
            flash(access.denied_message("experiment", exp.owner_username), "error")
            return redirect(page_url(exp))
        if exp.status != "planned":
            flash(gettext("“%(name)s” has already started.", name=exp.name), "error")
            return redirect(page_url(exp))
        exp.start_date = exp.start_date or date.today()
        exp.status = "active"
        exp.plan_json = json.dumps(plan.snapshot(s, exp, place), ensure_ascii=False)
        exp.updated_at = datetime.utcnow()
        s.commit()
        flash(gettext("Started %(name)s: day 1 is %(date)s. The plan is kept; a change to its groups, main readout, readout day or test will ask why.",  # noqa: E501
                      name=exp.name, date=i18n.strftime(exp.start_date, "%b %d, %Y")), "success")
        return redirect(page_url(exp))


@bp.post("/<int:experiment_id>/readings")
def save_readings(experiment_id: int):
    """{"on": date, "values": {subject key: "24.5" or "" (remove)}}: one
    cell from the sheet, or a whole day; `measure` for another measure's
    (its key). For a readout BioManager works out, each value is what was
    measured: {"length": "12.1", "width": "8.4"}."""
    data = request.get_json(silent=True) or {}
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        try:
            on = date.fromisoformat(str(data.get("on") or "")[:10])
        except ValueError:
            return jsonify({"ok": False, "error": gettext("Say which day it is for.")}), 400
        if on > date.today():
            return jsonify({"ok": False, "error": gettext("A readout can't be for a day still to come.")}), 400
        readout = readout_for(exp, place, data.get("measure"))
        if readout is None:
            return jsonify({"ok": False, "error": gettext("This experiment has no measure “%(measure)s”.",
                                                          measure=data.get("measure"))}), 400
        by_key = {x.key: x for x in subjects(s, exp, place)}
        problems = []
        given = list((data.get("values") or {}).items())
        for key, raw in control_first(readout, given, lambda item: by_key.get(item[0])):
            subject = by_key.get(key)
            if subject is None:
                continue
            problem = set_reading(s, exp, place, subject, on, raw if isinstance(raw, dict) else str(raw), readout)
            if problem:
                problems.append(problem)
        if problems and len(problems) == len(data.get("values") or {}):
            s.rollback()
            return jsonify({"ok": False, "error": " ".join(problems[:3])}), 400
        s.commit()
        s.refresh(exp)
        return jsonify({"ok": True, "problems": problems, **payload(s, exp, place)})


@bp.post("/<int:experiment_id>/measures/<int:measure_id>/update")
def update_measure(experiment_id: int, measure_id: int):
    """{"label", "unit", "control"}: another measure renamed, its unit or
    (for ΔΔCt and fold change) its control group changed. Its id stays, so
    its readings stay with it."""
    from . import audit
    data = request.get_json(silent=True) or request.form.to_dict()
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        before = readout_for(exp, place, f"{MEASURE}{measure_id}")
        try:
            with audit.batch(s, "update", f"Measure of {exp.name} changed", "experiments"):
                m = change_measure(exp, place, f"{MEASURE}{measure_id}", data)
                rekind_readings(s, exp, place, before, m)
        except LookupError:
            s.rollback()
            return jsonify({"ok": False, "error": gettext("This experiment has no measure “%(measure)s”.",
                                                          measure=measure_id)}), 404
        except (ValueError, derived.DerivedError) as e:
            s.rollback()
            return jsonify({"ok": False, "error": str(e)}), 409
        if m.get("per_group"):
            # A new control group: each day's ΔΔCt is against its animals now.
            subs = subjects(s, exp, place)
            s.flush()
            for on in set(s.scalars(select(ExperimentReading.read_on).where(
                    ExperimentReading.experiment_id_fk == exp.id, ExperimentReading.readout_key == m["key"]))):
                _store_day(s, exp, m, on, subs)
        exp.updated_at = datetime.utcnow()
        s.commit()
        s.refresh(exp)
        return jsonify({"ok": True, "measure": m["key"], **payload(s, exp, place)})


@bp.post("/<int:experiment_id>/measures/<int:measure_id>/delete")
def delete_measure(experiment_id: int, measure_id: int):
    """Another measure deleted with its readings, as one batch that Batch
    history's Undo brings back whole (the measure is in the experiment's
    readout JSON, so the experiment's own change is in the batch)."""
    from . import audit
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        refused = _refuse(exp)
        if refused:
            return refused
        m = readout_for(exp, place, f"{MEASURE}{measure_id}")
        if m is None or not m.get("measure"):
            return jsonify({"ok": False, "error": gettext("This experiment has no measure “%(measure)s”.",
                                                          measure=measure_id)}), 404
        rows = list(s.scalars(select(ExperimentReading).where(ExperimentReading.experiment_id_fk == exp.id,
                                                              ExperimentReading.readout_key == m["key"])))
        with audit.batch(s, "delete", f"Measure “{m['label']}” of {exp.name} deleted, with its {len(rows)} readings",
                         "experiments"):
            for row in rows:
                s.delete(row)
            drop_measure(exp, place, m["key"])
            exp.updated_at = datetime.utcnow()
        s.commit()
        s.refresh(exp)
        return jsonify({"ok": True, "message": ngettext("Deleted %(measure)s and its %(num)s reading. Batch history can undo it.",
                                                        "Deleted %(measure)s and its %(num)s readings. Batch history can undo it.",
                                                        len(rows), measure=m["label"]),
                        **payload(s, exp, place)})


@bp.post("/<int:experiment_id>/delete")
def delete(experiment_id: int):
    from .app import log_delete
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        if not access.can_edit_experiment(exp):
            flash(access.denied_message("experiment", exp.owner_username), "error")
            return redirect(page_url(exp))
        for model in (ExperimentSubject, ExperimentReading):
            for row in s.scalars(select(model).where(model.experiment_id_fk == exp.id)):
                s.delete(row)
        log_delete(s, "experiments", exp.id, record_label=f"Experiment: {exp.name}", details=place.key)
        name = exp.name
        s.delete(exp)
        s.commit()
    flash(gettext("Deleted experiment %(name)s. Its %(nouns)s are unchanged.", name=name,
                  nouns=i18n.translate_value(place.nouns)), "success")
    return redirect(place.list_url)


# ---------------------------------------------------------------- export

def _long_rows(exp, place, table, by_key) -> list[list]:
    """A readout's (or another measure's) values, one row per animal and day."""
    readout = table["readout"]
    unit = f" ({readout['unit']})" if readout["unit"] else ""
    measured = readout.get("inputs") or []      # a readout BioManager works out: what was measured, beside it
    long_rows = [["Experiment", "Database", place.noun.capitalize(), "Group", "Sex", "Genotype", place.housing.capitalize(),
                  "Date", "Day", f"{readout['label']}{unit}", "At start", "% of first" if readout["kind"] == "value" else "% of start",
                  "Left out of the analysis"] + [f"{i['label']} ({i['unit']})" if i["unit"] else i["label"] for i in measured]]
    for r in table["rows"]:
        x = by_key.get(r["key"])
        for i, d in enumerate(table["dates"]):
            if r["values"][i] is None:
                continue
            inputs = (r.get("inputs") or [None] * len(table["dates"]))[i] or {}
            long_rows.append([exp.name, place.label, r["label"], r["group"], r["sex"], x.genotype if x else "",
                              x.housing if x else "", d, table["days"][i], r["values"][i], r["start"], r["pct"][i],
                              r.get("left_out", "")] + [inputs.get(m["key"], "") for m in measured])
    return long_rows


def _stats_rows(table) -> list[list]:
    rows = []
    for day in table["stats"]["raw"]:
        rows.append([day["day"], day["date"],
                     "; ".join(f"{g['group']}: {g['value']}" + (f" ± {g['sem']}" if g.get("sem") is not None else "")
                               + f" (n={g['n']})" for g in day["groups"]),
                     day["test"], day["p"], day["stars"]])
    return rows


def _measure_exports(session, exp, place, subs) -> list[tuple[str, list[list], list[list]]]:
    """Each other measure with a value: (its name, its long rows, its tests by day)."""
    by_key = {x.key: x for x in subs}
    out = []
    for m in measures_of(exp):
        table = readout_table(session, exp, place, subs, m)
        if table["dates"]:
            out.append((m["label"], _long_rows(exp, place, table, by_key), _stats_rows(table)))
    return out


def _rows_for_export(session, exp, place):
    from . import experiment_steps as xs
    subs = subjects(session, exp, place)
    table = readout_table(session, exp, place, subs)
    by_key = {x.key: x for x in subs}
    long_rows = _long_rows(exp, place, table, by_key)
    wide = [[place.noun.capitalize(), "Group"] + [f"Day {n}" if n is not None else d for d, n in zip(table["dates"], table["days"])]]
    wide += [[r["label"], r["group"]] + r["values"] for r in table["rows"]]
    manip = [["Day", "Date planned", "Done on", "Done by", "What", "Kind", "Dose", "Route", "Group", place.noun.capitalize(),
              "Given", "Amount", "Volume", "Weight used", "Reagent", "Lot", "Note"]]
    for row in xs.schedule(session, exp, place, subs):
        if row["reading"]:
            continue
        step = next(st for st in exp.steps if st.id == row["step_id"])
        rec = row["record"]
        given = {e["subject"]: e for e in (rec["subjects"] if rec else [])}
        reagent = (rec or {}).get("reagent") or {}
        for x in subs:
            if step.treatment_group and x.group.strip().lower() != step.treatment_group.strip().lower():
                continue
            e = given.get(x.key)
            manip.append([row["day"], row["date"], rec["done_on"] if rec else "", rec["done_by"] if rec else "",
                          step.agent, kind_info(place.family, step.kind)[0], step.dose, step.route, step.treatment_group,
                          x.label, ("yes" if e else "no") if rec else "", (e or {}).get("amount", ""),
                          (e or {}).get("volume", ""), (e or {}).get("grams", ""), reagent.get("name", ""),
                          reagent.get("lot", ""), rec["note"] if rec else ""])
    stats = [["Day", "Date", "Groups", "Test", "p", ""]] + _stats_rows(table)
    return long_rows, wide, manip, stats


@bp.get("/<int:experiment_id>/export.<fmt>")
def export(experiment_id: int, fmt: str):
    """The readout (long, for Prism or R, and wide), each animal's
    manipulations, the day-by-day tests, each other measure on a sheet of
    its own and, for one started with a plan, the plan and the changes from
    it: an Excel workbook, or the long readout as CSV."""
    import csv
    import io
    import re as _re
    from flask import Response
    from . import experiment_plan
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        long_rows, wide, manip, stats = _rows_for_export(s, exp, place)
        measures = _measure_exports(s, exp, place, subjects(s, exp, place))
        plan_rows = experiment_plan.export_rows(s, exp, place)
        base = _re.sub(r"[^A-Za-z0-9._-]+", "-", exp.name).strip("-")[:60] or "experiment"
    from .services import csv_text
    if fmt == "csv":
        return Response(csv_text(long_rows), mimetype="text/csv",
                        headers={"Content-Disposition": f"attachment; filename={base}-readout.csv"})
    if fmt != "xlsx":
        abort(404)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    sheets = [("Readout", long_rows), ("Readout, wide", wide), ("Manipulations", manip), ("Tests by day", stats)]
    titles = {t.lower() for t, _rows in sheets} | {"plan and deviations"}
    for label, rows, _tests in measures:
        # A sheet's name: at most 31 characters, none of []:*?/\, and not another sheet's.
        title = base_title = _re.sub(r"[\[\]:*?/\\]+", " ", label).strip()[:31] or "Measure"
        n = 1
        while title.lower() in titles:
            n += 1
            title = f"{base_title[:27]} ({n})"
        titles.add(title.lower())
        sheets.append((title, rows))
    if plan_rows:
        sheets.append(("Plan and deviations", plan_rows))
    for i, (title, rows) in enumerate(sheets):
        sheet = book.active if i == 0 else book.create_sheet()
        sheet.title = title
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"          # someone's text, never a formula
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    tests = book["Tests by day"]
    for label, _rows, rows in measures:
        if rows:
            tests.append([])
            for row in [[label], *rows]:
                tests.append(row)
                for cell in tests[tests.max_row]:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.data_type = "s"
    if len(stats) > 1 or any(rows for _label, _long, rows in measures):
        tests.append([])
        tests.append(["Each day is tested on its own, not corrected for the number of days."])
    buffer = io.BytesIO()
    book.save(buffer)
    return Response(buffer.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={base}.xlsx"})


# ---------------------------------------------------------------- bench mode

@bp.get("/<int:experiment_id>/bench")
def bench(experiment_id: int):
    """One animal at a time, for a phone at the bench: weigh or count each,
    or give each today's manipulation and tick it."""
    with SessionLocal() as s:
        exp, place = _load(s, experiment_id)
        data = payload(s, exp, place)
        return render_template("experiment_bench.html", place=place, data=data, page_url=page_url(exp))


@bp.get("/<int:experiment_id>/bench.svg")
def bench_qr(experiment_id: int):
    """A QR code of bench mode's address, to open it on a phone."""
    from flask import Response
    from .labels import _qr_svg
    url = request.host_url.rstrip("/") + url_for("experiments.bench", experiment_id=experiment_id)
    svg = _qr_svg(url, scale=5)
    if not svg:
        abort(404)
    return Response(svg, mimetype="image/svg+xml", headers={"Cache-Control": "private, max-age=3600"})
